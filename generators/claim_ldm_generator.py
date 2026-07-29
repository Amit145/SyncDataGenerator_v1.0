import csv
import os
import shutil
from datetime import datetime, timedelta
from pathlib import Path

from helper.csv_writer import write_csv
from helper.raw_metadata import RAW_PULL_TS


CLAIMS_LDM_SCHEMAS = {'coverage.csv': ['coverage_identifier',
                  'coverage_type',
                  'coverage_description',
                  'coverage_start_date',
                  'coverage_end_date',
                  'coverage_type_identifier',
                  'maximum_deductible',
                  'limit_identifier',
                  'exclusion_identifier',
                  'deductible_identifier',
                  'benefit_identifier',
                  'risk_identifier',
                  'peril_identifier',
                  'product_element_identifier'],
 'policy_coverage.csv': ['policy_coverage_identifier',
                         'coverage_identifier',
                         'coverage_group_identifier',
                         'coverage_status_date',
                         'coverage_status_code',
                         'sum_insured',
                         'sum_insured_after_coinsurance',
                         'gross_annualized_premium',
                         'peril_premium',
                         'annual_aggregate_limit',
                         'excess',
                         'minimum_deductible',
                         'indexation',
                         'automatic_indexation_flag',
                         'coverage_extension',
                         'coverage_trigger',
                         'coverage_level_lh',
                         'logic_of_limit',
                         'scope_of_limit',
                         'pct_loss_deductible',
                         'actual_price',
                         'agreed_value',
                         'technical_price_natural_perils',
                         'mandatory_participation',
                         'liability_agreed_flag',
                         'liability_exposure_value',
                         'flag_dual_insurance',
                         'flag_inhabited_building',
                         'motor_accessories_sum_insured',
                         'time_element_sum_insured',
                         'value_insured_property_damage',
                         'contingent_bi_flag',
                         'contingent_bi_waiting_period',
                         'coverage_sum_at_risk',
                         'employer_contribution_level',
                         'premium_paid_previous_period',
                         'position_grade',
                         'benefit_grouping',
                         'insured_entity_identifier'],
 'policy.csv': ['policy_identifier',
                'policy_type',
                'policy_number',
                'aspire_policy_identifier',
                'policy_status_code',
                'policy_status_date',
                'policy_inception_date',
                'policy_inception_date_as_new_business',
                'policy_inception_date_of_first_allianz_policy',
                'policy_issue_date',
                'policy_end_date',
                'policy_expiry_date',
                'policy_cancellation_date',
                'policy_cancellation_notification_date',
                'policy_cancellation_reason',
                'policy_renewal_date',
                'policy_renewal_notification_date',
                'policy_duration',
                'premium',
                'gross_written_premium',
                'gross_earned_premium',
                'gross_original_premium',
                'paid_premium',
                'sum_of_paid_premium',
                'unearned_premium_amount',
                'adjustment_premium',
                'adjustment_earned_premium',
                'estimated_premium_income',
                'subject_premium_income',
                'technical_price',
                'technical_expected_loss',
                'projected_loss_ratio',
                'underwriting_premium_surcharge',
                'commission_due_amount',
                'gross_written_commission_amount',
                'gross_earned_commission',
                'gross_original_commission',
                'commission_share',
                'differed_acquisition_cost',
                'frequency_of_installments',
                'preferred_payment_date',
                'date_of_debit',
                'tax_code',
                'tariff_version',
                'premium_configuration_code',
                'ifrs17_measurement_model',
                'current_bonus_malus_class',
                'previous_bonus_malus_class',
                'no_claims_discount',
                'annual_aggregate_deductible',
                'limit_for_indemnification',
                'limit_business_interruption',
                'number_of_insured_persons',
                'number_of_substituting_policy',
                'automatic_renewal_years',
                'reinsurance_flag',
                'retention_flag',
                'tacit_renewal_flag',
                'multi_year_contracts',
                'lapse_flag',
                'flag_change_of_contract',
                'flag_cross_selling',
                'flag_quote_conversion',
                'marketing_campaign_flag',
                'primary_excess_liability_flag',
                'operational_paperless_consent',
                'legal_consent',
                'declared_driver',
                'broker_wording',
                'manual_clauses',
                'rationale_for_cession',
                'statute_of_limitation_date',
                'product_group_identifier',
                'policy_holder_identifier',
                'portfolio_identifier',
                'party_in_role_identifier',
                'policy_coverage_identifier',
                'coverage_identifier'],
 'physical_place.csv': ['physical_place_identifier',
                        'place_identifier',
                        'applicable_jurisdiction',
                        'census_zone',
                        'commune_code',
                        'geocoding_level',
                        'latitude',
                        'longitude',
                        'municipal_district',
                        'natcat_hazard_zone_scheme',
                        'point_of_interest',
                        'possible_max_business_interruption',
                        'sub_region',
                        'surface_elevation'],
 'loss_event.csv': ['loss_event_identifier',
                    'loss_event_name',
                    'loss_event_description',
                    'loss_date',
                    'loss_time',
                    'loss_event_start_date',
                    'loss_event_end_date',
                    'loss_event_status',
                    'loss_type',
                    'loss_category',
                    'loss_cause',
                    'property_liability_loss_cause',
                    'date_of_effective_loss',
                    'notification_date',
                    'notification_channel',
                    'days_from_event_to_fnol',
                    'time_to_notification',
                    'fatality_indicator',
                    'flag_catastrophe_event',
                    'flag_nat_cat_event',
                    'minimal_impact_flag',
                    'number_of_injured_parties',
                    'number_of_people_involved',
                    'number_of_vehicles_involved',
                    'claimant_drugs_or_alcohol',
                    'contributory_negligence',
                    'damage_specification',
                    'physical_place_identifier',
                    'claim_identifier'],
 'catastrophe.csv': ['catastrophe_identifier',
                     'catastrophe_cause',
                     'catastrophe_description',
                     'catastrophe_begin_date',
                     'catastrophe_end_date',
                     'cresta_zone',
                     'nat_cat_event_code',
                     'loss_event_identifier'],
 'person.csv': ['party_ref',
                'party_kind',
                'tenant_cd',
                'lead_ind',
                'paperless_ind',
                'src_party_ref',
                'src_party_type',
                'natural_ref',
                'given_nm',
                'family_nm',
                'display_nm',
                'title_txt',
                'occupation_txt',
                'dob',
                'birth_yr',
                'nationality_txt',
                'gender_txt',
                'marital_txt',
                'disability_degree',
                'language_pref',
                'role_txt',
                'job_title_txt',
                'legal_ref',
                'legal_name',
                'legal_score_no',
                'legal_status_txt',
                'legal_job_title_txt',
                'legal_src_ref',
                'legal_src_type',
                'constitution_dt',
                'lead_conv_dt',
                'email_home_txt',
                'email_work_txt',
                'phone_work_txt',
                'phone_home_txt',
                'person_identifier',
                'is_lead',
                'tenant_identifier',
                'person_type',
                'person_status',
                'preferred_language',
                'source_identifier',
                'source_type',
                'assessed_disability_degree',
                'is_operational_paperless_consent',
                'is_opt_in_legitimate_interest',
                'is_opt_in_validated',
                'digital_identifier',
                'identification_type',
                'legal_consent_flag',
                'credit_rating',
                'rate_class',
                'credit_rating_provider'],
 'claim_participant.csv': ['claim_participant_identifier',
                           'role_in_claim',
                           'involvement_type',
                           'role_start_date',
                           'role_end_date',
                           'injury_flag',
                           'is_primary_indicator',
                           'liability_percentage',
                           'person_identifier',
                           'claim_identifier'],
 'insured_entity.csv': ['insured_entity_identifier',
                        'entity_type',
                        'insured_person_identifier',
                        'insured_object_owner_identifier',
                        'insured_object_address_flag',
                        'insured_object_sum_insured',
                        'insured_object_description',
                        'insured_object_current_status',
                        'insured_object_start_date',
                        'insured_object_end_date',
                        'wall_construction_material_type',
                        'home_risk_address',
                        'home_type',
                        'home_state',
                        'vehicle_body_type',
                        'vehicle_fuel_type',
                        'vehicle_risk_address',
                        'vehicle_risk_class_code',
                        'vehicle_variant',
                        'vehicle_reg_state',
                        'vehicle_class',
                        'vehicle_model',
                        'exposure_base',
                        'grouped_location_flag',
                        'replacement_value',
                        'type_of_insured_entity',
                        'value_of_content',
                        'value_of_goods_carried',
                        'actual_pre_event_entity_value',
                        'damaged_entity'],
 'claim.csv': ['claim_identifier',
               'claim_type',
               'claim_number',
               'third_party_claim_number',
               'claim_status',
               'claim_state',
               'claim_status_date',
               'claim_open_date',
               'claim_close_date',
               'claim_reopen_date',
               'claim_reopen_reason',
               'claims_made_date',
               'movement_date',
               'claim_duration',
               'claim_type_code',
               'claim_process_method',
               'claim_sensitivity',
               'claim_specific_flag',
               'bodily_injury_indicator',
               'applicable_deductible_flag',
               'total_loss_flag',
               'litigation_flag',
               'proven_claim_fraud_flag',
               'hospitalized_indicator',
               'cross_border_claim_indicator',
               'intercompany_agreement_flag',
               'no_claims_discount',
               'claim_requested_amount',
               'outstanding_subrogation_amount',
               'recovery_actual',
               'recovery_type',
               'outstanding_reserve',
               'claims_expense',
               'is_claim_suspicious',
               'suspected_amt',
               'fraud_amt',
               'is_recovery_happened',
               'days_to_first_recovery',
               'days_to_last_recovery',
               'litigation_duration_days',
               'claim_reason',
               'claim_channel',
               'claim_product',
               'claim_band',
               'claim_band_sort',
               'claim_fraud_status',
               'claim_fraud_type',
               'claim_fraud_detection_method',
               'recovery_band',
               'recovery_category',
               'recovery_source',
               'claims_rejection_reason',
               'finalization_reason',
               'indemnity_logic',
               'coverage_verification_result',
               'instruction_closure_date',
               'subrogation_paid_date',
               'claims_history_lob',
               'pre_accident_work_status',
               'return_to_work_status',
               'fire_brigade_fees_and_levies',
               'claims_external_fte_count',
               'total_incurred',
               'total_payment_amount',
               'reimbursable_claim_amount',
               'reimbursable_treatment_amount',
               'settlement_amount_pc',
               'claim_approval_date',
               'claim_payment_date',
               'claim_amounts_description',
               'policy_coverage_identifier',
               'policy_identifier',
               'insured_entity_identifier',
               'coverage_identifier'],
 'police_report.csv': ['police_report_identifier',
                       'police_report_number',
                       'police_report_date',
                       'police_report_flag',
                       'claim_identifier'],
 'medical_report.csv': ['medical_report_identifier',
                        'document_identifier',
                        'health_declaration',
                        'claim_identifier'],
 'treatment.csv': ['treatment_identifier',
                   'treatment_code',
                   'treatment_code_standard',
                   'treatment_description',
                   'treatment_complexity_level',
                   'initial_treatment_date',
                   'treatment_start_date',
                   'number_of_treatments',
                   'paid_treatment_amount',
                   'treatment_amount_currency',
                   'medication_code',
                   'treatment_type',
                   'claim_event_identifier',
                   'medical_condition_identifier',
                   'claim_participant_identifier'],
 'medical_condition.csv': ['medical_condition_identifier',
                           'disease_name',
                           'disability_status',
                           'pre_existing_condition'],
 'health_insurance_claim.csv': ['health_insurance_claim_identifier',
                                'claim_identifier',
                                'date_of_first_medical_expert_report',
                                'health_data_consent'],
 'claim_event.csv': ['claim_event_identifier',
                     'event_type',
                     'event_date',
                     'event_description',
                     'claim_identifier'],
 'litigation.csv': ['litigation_identifier',
                    'claim_event_identifier',
                    'court_matter_type',
                    'date_of_legal_representation',
                    'first_litigation_date',
                    'litigation_date',
                    'decision_of_court',
                    'claim_identifier'],
 'settlement.csv': ['settlement_identifier',
                    'claim_event_identifier',
                    'date_of_final_settlement',
                    'date_of_settlement_offer_decision',
                    'insurers_current_settlement_offer_amount',
                    'insurers_first_settlement_offer_amount',
                    'outstanding_medical_expenses',
                    'settlement_table',
                    'settlement_type'],
 'medical_assessment.csv': ['medical_assessment_identifier',
                            'claim_event_identifier',
                            'assessed_disability_degree',
                            'permanent_disability_degree',
                            'permanent_disability_flag',
                            'medical_report_flag',
                            'medical_report_date',
                            'person_hospital_status',
                            'psychological_factors',
                            'medical_condition_description'],
 'claim_investigation.csv': ['claim_investigation_identifier',
                             'claim_event_identifier',
                             'claim_handler_identifier',
                             'claim_handler_notes',
                             'claim_investigation_start_date',
                             'claim_investigation_end_date',
                             'fraud_indicator',
                             'investigator_flag',
                             'claim_fraud_assessment'],
 'repair.csv': ['repair_identifier',
                'claim_event_identifier',
                'number_of_spare_parts_repaired',
                'paid_spare_parts_costs',
                'spare_part_repair_shop_price',
                'spare_part_type',
                'claim_participant_identifier'],
 'injury.csv': ['injury_identifier', 'claim_event_identifier', 'bodily_injury_date', 'injury_description'],
 'diagnosis.csv': ['diagnosis_identifier',
                   'diagnosis_code',
                   'diagnosis_description',
                   'diagnosis_priority',
                   'diagnosis_standard',
                   'medical_condition_identifier',
                   'injury_identifier',
                   'claim_event_identifier']}


CLAIMS_METADATA_COLUMNS = ["batch_ref", "pull_ts", "origin_sys"]

CLAIMS_LDM_SCHEMAS = {
    name: CLAIMS_METADATA_COLUMNS + [column for column in schema if column not in CLAIMS_METADATA_COLUMNS]
    for name, schema in CLAIMS_LDM_SCHEMAS.items()
}


CLAIMS_RAW_FILE_NAMES = {
    "coverage.csv": "coverage_catalog.csv",
    "policy_coverage.csv": "policy_coverage_register.csv",
    "person.csv": "party_master.csv",
    "claim_participant.csv": "claim_participant_register.csv",
    "insured_entity.csv": "insured_entity_register.csv",
    "policy.csv": "policy_register.csv",
    "claim.csv": "claim_register.csv",
    "loss_event.csv": "loss_event_register.csv",
    "claim_event.csv": "claim_event_log.csv",
    "physical_place.csv": "physical_place_register.csv",
    "catastrophe.csv": "catastrophe_event_register.csv",
    "treatment.csv": "treatment_register.csv",
    "health_insurance_claim.csv": "health_claim_register.csv",
    "police_report.csv": "police_report_register.csv",
    "medical_report.csv": "medical_report_register.csv",
    "settlement.csv": "settlement_register.csv",
    "injury.csv": "injury_register.csv",
    "claim_investigation.csv": "claim_investigation_register.csv",
    "repair.csv": "repair_register.csv",
    "litigation.csv": "litigation_register.csv",
    "medical_assessment.csv": "medical_assessment_register.csv",
    "medical_condition.csv": "medical_condition_catalog.csv",
    "diagnosis.csv": "diagnosis_register.csv",
}

CLAIMS_TWO_SOURCE_SPEC_PATH = Path(__file__).resolve().parents[1] / "claims" / "Claims_2Sources_DataTables.xlsx"
CLAIMS_TWO_SOURCE_TABLES = {
    "Claim": ("claim.csv", "claim_register.csv"),
    "Loss Event": ("loss_event.csv", "loss_event_register.csv"),
    "Claim Investigation": ("claim_investigation.csv", "claim_investigation_register.csv"),
}

CLAIMS_PRD2_CLAIM_EXTRA_COLUMNS = [
    "outstanding_reserve",
    "claims_expense",
    "is_claim_suspicious",
    "suspected_amt",
    "fraud_amt",
    "is_recovery_happened",
    "days_to_first_recovery",
    "days_to_last_recovery",
    "litigation_duration_days",
    "claim_reason",
    "claim_channel",
    "claim_product",
    "claim_band",
    "claim_band_sort",
    "claim_fraud_status",
    "claim_fraud_type",
    "claim_fraud_detection_method",
    "recovery_band",
    "recovery_category",
    "recovery_source",
]


def _snake(value: str) -> str:
    return "_".join(str(value or "").strip().lower().replace("-", " ").split())


def _first_value(*values) -> str:
    for value in values:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def _batch_pull_ts(batch_id: str) -> str:
    return RAW_PULL_TS


def _read_csv_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _copy_csv(src: Path, dst_dir: Path, dst_name: str | None = None) -> None:
    dst_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dst_dir / (dst_name or src.name))


def _load_claims_two_source_spec() -> dict[str, list[dict]]:
    """Read the claims source-1/source-2 workbook into raw column mappings."""
    if not CLAIMS_TWO_SOURCE_SPEC_PATH.exists():
        raise FileNotFoundError(f"Claims two-source spec not found: {CLAIMS_TWO_SOURCE_SPEC_PATH}")
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("openpyxl is required to read claims/Claims_2Sources_DataTables.xlsx") from exc

    wb = openpyxl.load_workbook(CLAIMS_TWO_SOURCE_SPEC_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]
    mapping: dict[str, list[dict]] = {raw_table: [] for raw_table, _ in CLAIMS_TWO_SOURCE_TABLES.values()}
    for values in ws.iter_rows(min_row=2, values_only=True):
        table_label = str(values[1] or "").strip()
        logical_attr = str(values[2] or "").strip()
        src1_attr = str(values[4] or "").strip()
        src2_attr = str(values[6] or "").strip()
        if not table_label or not logical_attr or table_label not in CLAIMS_TWO_SOURCE_TABLES:
            continue
        raw_table, _ = CLAIMS_TWO_SOURCE_TABLES[table_label]
        logical_col = _snake(logical_attr)
        if logical_col not in CLAIMS_LDM_SCHEMAS[raw_table]:
            continue
        mapping[raw_table].append({
            "logical": logical_col,
            "src1": _snake(src1_attr) if src1_attr else "",
            "src2": _snake(src2_attr) if src2_attr else "",
        })
    claim_mapping = mapping["claim.csv"]
    existing_src2 = {entry["src2"] for entry in claim_mapping}
    for column in CLAIMS_PRD2_CLAIM_EXTRA_COLUMNS:
        if column not in existing_src2:
            claim_mapping.append({"logical": column, "src1": column, "src2": column})
    return mapping


def _write_claims_source_view(
    source_dir: Path,
    table_name: str,
    rows: list[dict],
    mapping_rows: list[dict],
    batch_id: str,
    origin_sys: str,
) -> None:
    columns = ["batch_ref", "pull_ts", "origin_sys"]
    columns.extend(entry["src2"] for entry in mapping_rows if entry["src2"])
    out_rows = []
    extract_ts = _batch_pull_ts(batch_id)
    for row_index, row in enumerate(rows):
        out = {
            "batch_ref": batch_id,
            "pull_ts": extract_ts,
            "origin_sys": origin_sys,
        }
        for entry in mapping_rows:
            source_col = entry["src2"]
            if source_col:
                value = row.get(entry["logical"], "")
                if table_name == "claim.csv" and source_col == "recovery_category" and row.get("recovery_category"):
                    value = row.get("recovery_category", "")
                if table_name == "claim.csv" and source_col == "claim_ref_id" and row_index % 3 == 0:
                    value = f"{origin_sys}_{value}"
                out[source_col] = value
        out_rows.append(out)
    write_csv(str(source_dir), table_name, out_rows, fieldnames=columns)


def _rebuild_claims_table_from_sources(
    src1_rows: list[dict],
    src2_rows: list[dict],
    mapping_rows: list[dict],
    schema: list[str],
) -> list[dict]:
    rebuilt = []
    for index, src1 in enumerate(src1_rows):
        src2 = src2_rows[index] if index < len(src2_rows) else {}
        row = {column: "" for column in schema}
        for column in CLAIMS_METADATA_COLUMNS:
            if column in row:
                row[column] = src1.get(column, "") or src2.get(column, "")
        for entry in mapping_rows:
            value = src1.get(entry["logical"], "")
            if str(value).strip() == "" and entry["src2"]:
                value = src2.get(entry["src2"], "")
            row[entry["logical"]] = value
        rebuilt.append(row)
    return rebuilt


def write_claims_two_source_raw(claims_raw_dir: str, raw_root: str, batch_id: str) -> dict[str, str]:
    """Create claims PRD1/PRD2 source views and a vault-ready consolidated raw folder.

    PRD1 is the complete LDM raw source. PRD2 follows the SAP/source-2
    attribute names from claims/Claims_2Sources_DataTables.xlsx for the tables
    defined in that workbook. The consolidated raw_vault folder is shaped back
    to the existing claims LDM raw schema so the current bronze/silver/gold
    claims pipeline can load it without logic changes.
    """
    source_raw = Path(claims_raw_dir)
    root = Path(raw_root) / "claims" / batch_id
    prd1_dir = root / "prd_01"
    prd2_dir = root / "prd_02"
    raw_vault_dir = root / "raw_vault"
    for path in (prd1_dir, prd2_dir, raw_vault_dir):
        if path.exists():
            shutil.rmtree(path)
        path.mkdir(parents=True, exist_ok=True)

    mappings = _load_claims_two_source_spec()

    for raw_file, raw_name in CLAIMS_RAW_FILE_NAMES.items():
        source_path = source_raw / raw_name
        if source_path.exists():
            _copy_csv(source_path, prd1_dir, raw_name)
            _copy_csv(source_path, raw_vault_dir, raw_name)

    for raw_file, raw_name in CLAIMS_RAW_FILE_NAMES.items():
        if raw_file not in mappings:
            continue
        source_path = source_raw / raw_name
        rows = _read_csv_rows(source_path)
        mapping_rows = mappings[raw_file]
        prd_file = raw_file
        _write_claims_source_view(prd2_dir, prd_file, rows, mapping_rows, batch_id, "SAP")
        src1_rows = _read_csv_rows(prd1_dir / raw_name)
        src2_rows = _read_csv_rows(prd2_dir / prd_file)
        rebuilt = _rebuild_claims_table_from_sources(src1_rows, src2_rows, mapping_rows, CLAIMS_LDM_SCHEMAS[raw_file])
        write_csv(str(raw_vault_dir), raw_name, rebuilt, fieldnames=CLAIMS_LDM_SCHEMAS[raw_file])

    with (root / "_source_manifest.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["source", "folder", "description"])
        writer.writeheader()
        writer.writerows([
            {"source": "prd_01", "folder": str(prd1_dir), "description": "Claims LDM raw source-1 files"},
            {"source": "prd_02", "folder": str(prd2_dir), "description": "Claims SAP/source-2 view from Claims_2Sources_DataTables.xlsx"},
            {"source": "raw_vault", "folder": str(raw_vault_dir), "description": "Vault-ready consolidated claims raw in existing LDM schema"},
        ])

    return {
        "prd_01": str(prd1_dir),
        "prd_02": str(prd2_dir),
        "raw_vault": str(raw_vault_dir),
    }


def _as_list(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def _index_by(rows, key):
    return {row[key]: row for row in rows}


def _invert_multi_map(mapping):
    out = {}
    for left, right_values in mapping.items():
        for right in _as_list(right_values):
            out[right] = left
    return out


def _date(value, default):
    if not value:
        return default
    return str(value).split("T")[0].split(" ")[0]


def _iso(day):
    if isinstance(day, str):
        day = datetime.fromisoformat(day[:10])
    return day.strftime("%Y-%m-%dT00:00:00")


def _time(day):
    if isinstance(day, str):
        day = datetime.fromisoformat(day[:19])
    return day.strftime("%H:%M:%S")


def _amount(value, default):
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return default


def _bounded_date(day, start_day, end_day):
    if end_day < start_day:
        end_day = start_day
    return max(start_day, min(day, end_day))


def _source_identifier(prefix, value):
    value = str(value or "").strip()
    return f"{prefix}_{value}" if value else ""


def _has_existing(rows, key, value):
    return any(row.get(key) == value for row in rows)


def _jurisdiction_from_country(country):
    normalized = str(country or "").strip().upper()
    if normalized in {"UK", "GB", "GBR", "UNITED KINGDOM", "ENGLAND", "SCOTLAND", "WALES", "NORTHERN IRELAND"}:
        return "UK"
    if normalized in {"IN", "IND", "INDIA"}:
        return "India"
    return str(country or "").strip() or "UK"


def write_claims_ldm_raw_batch(base_folder, batch_id, ctx):
    out_dir = os.path.join(base_folder, "claims", batch_id, "raw")
    os.makedirs(out_dir, exist_ok=True)

    hub_person_by_hk = _index_by(ctx["hub_person_rows"], "Person Hash Key")
    hub_contact_by_hk = _index_by(ctx["hub_con_rows"], "Contact Hash Key")
    hub_nat_by_hk = _index_by(ctx["hub_nat"], "Natural Person Hash Key")
    hub_leg_by_hk = _index_by(ctx["hub_leg"], "Legal Person Hash Key")
    hub_policy_by_hk = _index_by(ctx["hub_pol_rows"], "Policy Hash Key")
    hub_product_by_hk = _index_by(ctx["hub_prod_rows"], "Product Hash Key")

    sat_person_by_hk = _index_by(ctx["sat_per"], "Person Hash Key")
    sat_contact_by_hk = _index_by(ctx["sat_con"], "Contact Hash Key")
    sat_nat_by_hk = _index_by(ctx["sat_nat"], "Natural Person Hash Key")
    sat_leg_by_hk = _index_by(ctx["sat_leg"], "Legal Person Hash Key")
    sat_policy_by_hk = _index_by(ctx["sat_pol"], "Policy Hash Key")
    sat_motor_by_hk = _index_by(ctx["sat_mot"], "Motor Hash Key")
    sat_home_by_hk = _index_by(ctx["sat_hom"], "Home Hash Key")

    person_by_policy_hk = _invert_multi_map(ctx["policy_person_map"])
    contact_hk_by_person = {
        person_hk: _as_list(contact_hks)[0]
        for person_hk, contact_hks in ctx["person_to_contact"].items()
        if _as_list(contact_hks)
    }
    nat_by_person_hk = {
        person_hk: _as_list(nat_hks)[0]
        for person_hk, nat_hks in ctx["person_to_nat"].items()
        if _as_list(nat_hks)
    }
    leg_by_person_hk = dict(ctx["person_to_leg"])
    product_hk_by_code = {code: hk for hk, code in ctx["product_code_by_hk"].items()}

    rows = {name: [] for name in CLAIMS_LDM_SCHEMAS}
    applicable_jurisdiction = _jurisdiction_from_country(ctx.get("country", "UK"))
    place_templates = [
        ("Manchester", "53.4808", "-2.2426", "Medium", "Urban Junction", "38", "260000", "CEN004", "COM004", "ROOFTOP", "North West"),
        ("London", "51.5074", "-0.1278", "Low", "Urban Street", "11", "420000", "CEN005", "COM005", "ROOFTOP", "Greater London"),
        ("Birmingham", "52.4862", "-1.8904", "Medium", "Residential Area", "140", "210000", "CEN006", "COM006", "STREET", "West Midlands"),
        ("Glasgow", "55.8642", "-4.2518", "Medium", "Commercial Zone", "40", "230000", "CEN007", "COM007", "PARCEL", "Scotland"),
    ]
    condition_templates = [
        ("MC_FRAC", "Fracture", "Temporary", "N"),
        ("MC_BURN", "Burn Injury", "Temporary", "N"),
        ("MC_BACK", "Chronic Back Pain", "Permanent", "Y"),
        ("MC_MINOR", "Minor External Injury", "", "N"),
    ]
    for condition in condition_templates:
        rows["medical_condition.csv"].append({
            "medical_condition_identifier": condition[0],
            "disease_name": condition[1],
            "disability_status": condition[2],
            "pre_existing_condition": condition[3],
        })

    seen_people = set()
    claim_seq = 1
    event_seq = 1
    participant_seq = 1
    for policy_index, (policy_hk, hub_policy) in enumerate(hub_policy_by_hk.items(), start=1):
        person_hk = person_by_policy_hk.get(policy_hk)
        if not person_hk:
            continue
        sat_policy = sat_policy_by_hk.get(policy_hk, {})
        person = hub_person_by_hk.get(person_hk, {})
        sat_person = sat_person_by_hk.get(person_hk, {})
        nat_hk = nat_by_person_hk.get(person_hk)
        leg_hk = leg_by_person_hk.get(person_hk)
        nat = sat_nat_by_hk.get(nat_hk, {}) if nat_hk else {}
        leg = sat_leg_by_hk.get(leg_hk, {}) if leg_hk else {}
        contact_hk = contact_hk_by_person.get(person_hk)
        sat_contact = sat_contact_by_hk.get(contact_hk, {}) if contact_hk else {}

        person_identifier = _source_identifier("PER", person.get("Person Id"))
        if person_identifier and person_identifier not in seen_people:
            seen_people.add(person_identifier)
            person_type = sat_person.get("Type", "NATURAL")
            credit_rating = "A" if policy_index % 3 else "B"
            if person_type == "LEGAL":
                rate_class = "Commercial"
            elif credit_rating == "A" and sat_person.get("Operational Paperless Consent", "Y") == "Y":
                rate_class = "Preferred"
            else:
                rate_class = "Standard"
            rows["person.csv"].append({
                "party_ref": person.get("Person Id"),
                "party_kind": person_type,
                "tenant_cd": sat_person.get("Tenant Id", f"TENANT_{policy_index % 5 + 1:03d}"),
                "lead_ind": sat_person.get("Is Lead", "N"),
                "paperless_ind": sat_person.get("Operational Paperless Consent", "Y"),
                "src_party_ref": person_identifier,
                "src_party_type": "CLAIMS",
                "natural_ref": hub_nat_by_hk.get(nat_hk, {}).get("Natural Person Id", "") if nat_hk else "",
                "given_nm": nat.get("First Name", "") if person_type != "LEGAL" else "",
                "family_nm": nat.get("Last Name", "") if person_type != "LEGAL" else "",
                "display_nm": nat.get("Full Name", "") if person_type != "LEGAL" else leg.get("Company Name", ""),
                "title_txt": nat.get("Courtesy Title", "") if person_type != "LEGAL" else "",
                "occupation_txt": nat.get("Occupation", "") if person_type != "LEGAL" else "",
                "dob": nat.get("Birth Date", "") if person_type != "LEGAL" else "",
                "birth_yr": nat.get("Birth Year", "") if person_type != "LEGAL" else "",
                "nationality_txt": nat.get("Nationality", "") if person_type != "LEGAL" else "",
                "gender_txt": nat.get("Gender", "") if person_type != "LEGAL" else "",
                "marital_txt": nat.get("Marital Status", "") if person_type != "LEGAL" else "",
                "disability_degree": nat.get("Assesed Disability Degree", "") if person_type != "LEGAL" else "",
                "language_pref": sat_person.get("Preferred Language", "English"),
                "role_txt": nat.get("Role", "") if person_type != "LEGAL" else "Policy Holder",
                "job_title_txt": nat.get("Job Title", "") if person_type != "LEGAL" else leg.get("Job Title", ""),
                "legal_ref": hub_leg_by_hk.get(leg_hk, {}).get("Legal Person Id", "") if leg_hk else "",
                "legal_name": leg.get("Company Name", "") if person_type == "LEGAL" else "",
                "legal_score_no": leg.get("Person Score", "") if person_type == "LEGAL" else "",
                "legal_status_txt": leg.get("Person Status", "") if person_type == "LEGAL" else "",
                "legal_job_title_txt": leg.get("Job Title", "") if person_type == "LEGAL" else "",
                "legal_src_ref": leg.get("Source Id", "") if person_type == "LEGAL" else "",
                "legal_src_type": leg.get("Source Type", "") if person_type == "LEGAL" else "",
                "constitution_dt": _date(leg.get("Date of Constitution", ""), "") if person_type == "LEGAL" else "",
                "lead_conv_dt": leg.get("Converted Date", "") if person_type == "LEGAL" else "",
                "email_home_txt": sat_contact.get("Personal Email", ""),
                "email_work_txt": sat_contact.get("Work Email", ""),
                "phone_work_txt": sat_contact.get("Work Phone", ""),
                "phone_home_txt": sat_contact.get("Home Phone", ""),
                "person_identifier": person_identifier,
                "person_type": person_type,
                "person_status": "Active",
                "tenant_identifier": f"TENANT_{policy_index % 5 + 1:03d}",
                "source_identifier": person_identifier,
                "source_type": "CLAIMS",
                "preferred_language": sat_person.get("Preferred Language", "English"),
                "credit_rating": credit_rating,
                "credit_rating_provider": "Experian",
                "rate_class": rate_class,
                "identification_type": "Driving Licence" if nat else "Company Registration",
                "digital_identifier": f"DIGI{policy_index:08d}",
                "is_lead": sat_person.get("Is Lead", "N"),
                "is_operational_paperless_consent": sat_person.get("Operational Paperless Consent", "Y"),
                "is_opt_in_legitimate_interest": sat_person.get("Legal Consent", "Y"),
                "is_opt_in_validated": "Y",
                "legal_consent_flag": sat_person.get("Legal Consent", "Y"),
                "assessed_disability_degree": "0.00",
            })

        policy_identifier = _source_identifier("POL", hub_policy.get("Policy Id"))
        product_code = ctx["policy_to_product_id"].get(policy_hk, "")
        product_hk = product_hk_by_code.get(product_code)
        product_identifier = _source_identifier("PROD", hub_product_by_hk.get(product_hk, {}).get("Product Id") or product_code)
        policy_start = datetime.fromisoformat(_date(sat_policy.get("Policy Start Date"), "2024-01-01"))
        policy_end = datetime.fromisoformat(_date(sat_policy.get("Policy End Date"), "2024-12-31"))
        premium = _amount(sat_policy.get("Renewal Amount Current Period"), 900.0)
        policy_type = "Motor" if ctx["policy_to_motor"].get(policy_hk) else "Property" if ctx["policy_to_home"].get(policy_hk) else "Liability"
        policy_issue_date = _bounded_date(policy_start - timedelta(days=7), policy_start - timedelta(days=30), policy_start)
        policy_duration = max((policy_end - policy_start).days, 0)
        policy_renewal_date = _bounded_date(policy_end - timedelta(days=30), policy_start, policy_end)
        policy_renewal_notice_date = _bounded_date(policy_renewal_date - timedelta(days=14), policy_start, policy_renewal_date)
        policy_status = sat_policy.get("Policy Status", "ACTIVE")
        is_cancelled = str(policy_status or "").upper() in {"CANCELLED", "CANCELED", "LAPSED"}
        replacement_value_for_policy = round(max(premium * 80, 50000), 2)
        payment_date = _bounded_date(policy_start + timedelta(days=min(5, policy_duration)), policy_start, policy_end)
        debit_date = _bounded_date(payment_date + timedelta(days=2), policy_start, policy_end)
        coverage_identifier = f"COV_{policy_type.upper()}_{(policy_index % 5) + 1:03d}"
        policy_coverage_identifier = f"PC_{batch_id}_{policy_index:06d}"
        rows["policy.csv"].append({
            "policy_identifier": policy_identifier,
            "policy_number": hub_policy.get("Policy Number", policy_identifier),
            "policy_type": policy_type,
            "aspire_policy_identifier": f"ASP_{policy_identifier}",
            "policy_status_code": policy_status,
            "policy_status_date": _iso(policy_start),
            "policy_inception_date": _iso(policy_start),
            "policy_inception_date_as_new_business": _iso(policy_start),
            "policy_inception_date_of_first_allianz_policy": _iso(policy_start),
            "policy_issue_date": _iso(policy_issue_date),
            "policy_end_date": _iso(policy_end),
            "policy_expiry_date": _iso(policy_end),
            "policy_cancellation_date": _iso(policy_end) if is_cancelled else "",
            "policy_cancellation_notification_date": _iso(policy_renewal_notice_date) if is_cancelled else "",
            "policy_cancellation_reason": "Non Payment" if is_cancelled else "",
            "policy_renewal_date": _iso(policy_renewal_date),
            "policy_renewal_notification_date": _iso(policy_renewal_notice_date),
            "policy_duration": policy_duration,
            "premium": premium,
            "gross_written_premium": premium,
            "gross_earned_premium": round(premium * 0.92, 2),
            "gross_original_premium": round(premium * 0.97, 2),
            "paid_premium": round(premium * 0.88, 2),
            "sum_of_paid_premium": round(premium * 0.88, 2),
            "unearned_premium_amount": round(premium * 0.08, 2),
            "adjustment_premium": round(premium * 0.02, 2),
            "adjustment_earned_premium": round(premium * 0.01, 2),
            "estimated_premium_income": round(premium * 1.04, 2),
            "subject_premium_income": premium,
            "technical_price": round(premium * 0.82, 2),
            "technical_expected_loss": round(premium * 0.55, 2),
            "projected_loss_ratio": 0.55,
            "underwriting_premium_surcharge": round(premium * 0.03, 2),
            "commission_due_amount": round(premium * 0.12, 2),
            "gross_written_commission_amount": round(premium * 0.12, 2),
            "gross_earned_commission": round(premium * 0.10, 2),
            "gross_original_commission": round(premium * 0.11, 2),
            "commission_share": 0.12,
            "differed_acquisition_cost": round(premium * 0.04, 2),
            "frequency_of_installments": "Monthly" if policy_index % 2 else "Annual",
            "preferred_payment_date": _iso(payment_date),
            "date_of_debit": _iso(debit_date),
            "tax_code": "IPT",
            "tariff_version": f"TV{policy_start.year}",
            "premium_configuration_code": f"PCFG_{policy_type.upper()}",
            "ifrs17_measurement_model": "PAA",
            "current_bonus_malus_class": f"BM{policy_index % 10 + 1}",
            "previous_bonus_malus_class": f"BM{policy_index % 10}",
            "no_claims_discount": max(0, 30 - (policy_index % 6) * 5),
            "annual_aggregate_deductible": round(max(premium * 0.20, 250), 2),
            "limit_for_indemnification": round(replacement_value_for_policy * 1.25, 2),
            "limit_business_interruption": round(replacement_value_for_policy * 0.15, 2) if policy_type == "Property" else 0,
            "number_of_insured_persons": 1,
            "number_of_substituting_policy": policy_index % 2,
            "automatic_renewal_years": 1,
            "reinsurance_flag": "Y" if replacement_value_for_policy > 100000 else "N",
            "retention_flag": "Y",
            "tacit_renewal_flag": "Y",
            "multi_year_contracts": "N",
            "lapse_flag": "Y" if str(policy_status or "").upper() == "LAPSED" else "N",
            "flag_change_of_contract": "N",
            "flag_cross_selling": "Y" if policy_index % 4 == 0 else "N",
            "flag_quote_conversion": "Y",
            "marketing_campaign_flag": "Y" if policy_index % 5 == 0 else "N",
            "primary_excess_liability_flag": "Y" if policy_type == "Liability" else "N",
            "operational_paperless_consent": sat_person.get("Operational Paperless Consent", "Y"),
            "legal_consent": sat_person.get("Legal Consent", "Y"),
            "declared_driver": person_identifier if policy_type == "Motor" else "",
            "broker_wording": f"{policy_type} standard wording",
            "manual_clauses": "Standard exclusions apply",
            "rationale_for_cession": "Risk retention within treaty limits",
            "statute_of_limitation_date": _iso(policy_end + timedelta(days=365 * 3)),
            "product_group_identifier": f"PG_{policy_type.upper()}",
            "policy_holder_identifier": person_identifier,
            "portfolio_identifier": f"PORT_{policy_type.upper()}_{policy_index % 5 + 1:03d}",
            "party_in_role_identifier": f"PIR_{batch_id}_{policy_index:06d}",
            "policy_coverage_identifier": policy_coverage_identifier,
            "coverage_identifier": coverage_identifier,
        })

        if not _has_existing(rows["coverage.csv"], "coverage_identifier", coverage_identifier):
            rows["coverage.csv"].append({
                "coverage_identifier": coverage_identifier,
                "coverage_code": "COMP" if policy_type == "Motor" else "FIRE" if policy_type == "Property" else "LIAB",
                "coverage_type": policy_type,
                "coverage_description": f"{policy_type} comprehensive coverage",
                "coverage_start_date": _iso(policy_start),
                "coverage_end_date": _iso(policy_end),
                "maximum_deductible": round(max(premium * 8, 1000), 2),
            })

        insured_entity_identifier = f"IE_{batch_id}_{policy_index:06d}"
        insured_entity_type = "Vehicle" if policy_type == "Motor" else "Property" if policy_type == "Property" else "Liability"
        replacement_value = round(max(premium * 80, 50000), 2)
        motor_hk = ctx["policy_to_motor"].get(policy_hk)
        home_hk = ctx["policy_to_home"].get(policy_hk)
        sat_motor = sat_motor_by_hk.get(motor_hk, {}) if motor_hk else {}
        sat_home = sat_home_by_hk.get(home_hk, {}) if home_hk else {}
        if policy_type == "Motor":
            insured_object_description = " ".join(
                part for part in [
                    str(sat_motor.get("Vehicle Year", "")).strip(),
                    str(sat_motor.get("Vehicle Model", "")).strip(),
                    str(sat_motor.get("Variant", "")).strip(),
                ]
                if part
            )
        elif policy_type == "Property":
            insured_object_description = " ".join(
                part for part in [
                    str(sat_home.get("Home Type", "")).strip(),
                    str(sat_home.get("Home Risk Address", "")).strip(),
                ]
                if part
            )
        else:
            insured_object_description = "Third party liability exposure"
        rows["insured_entity.csv"].append({
            "insured_entity_identifier": insured_entity_identifier,
            "entity_type": insured_entity_type,
            "insured_person_identifier": person_identifier,
            "insured_object_owner_identifier": person_identifier,
            "insured_object_address_flag": "Y" if policy_type in {"Motor", "Property"} else "N",
            "insured_object_sum_insured": replacement_value,
            "insured_object_description": insured_object_description,
            "insured_object_current_status": policy_status,
            "insured_object_start_date": _iso(policy_start),
            "insured_object_end_date": _iso(policy_end),
            "wall_construction_material_type": sat_home.get("Wall Construction", "") if policy_type == "Property" else "",
            "home_risk_address": sat_home.get("Home Risk Address", "") if policy_type == "Property" else "",
            "home_type": sat_home.get("Home Type", "") if policy_type == "Property" else "",
            "home_state": sat_home.get("Home State", "") if policy_type == "Property" else "",
            "vehicle_body_type": sat_motor.get("Body Type", "") if policy_type == "Motor" else "",
            "vehicle_fuel_type": sat_motor.get("Fuel Type", "") if policy_type == "Motor" else "",
            "vehicle_risk_address": sat_motor.get("Motor Risk Address", "") if policy_type == "Motor" else "",
            "vehicle_risk_class_code": sat_motor.get("Risk Class Code", "") if policy_type == "Motor" else "",
            "vehicle_variant": sat_motor.get("Variant", "") if policy_type == "Motor" else "",
            "vehicle_reg_state": sat_motor.get("Vehicle RegState", "") if policy_type == "Motor" else "",
            "vehicle_class": sat_motor.get("Vehicle Class", "") if policy_type == "Motor" else "",
            "vehicle_model": sat_motor.get("Vehicle Model", "") if policy_type == "Motor" else "",
            "exposure_base": "IDV" if policy_type == "Motor" else "Reinstatement",
            "replacement_value": replacement_value,
            "type_of_insured_entity": insured_entity_type,
            "value_of_content": round(replacement_value * 0.08, 2),
            "value_of_goods_carried": round(replacement_value * 0.03, 2) if policy_type == "Motor" else "",
            "actual_pre_event_entity_value": round(replacement_value * 0.92, 2),
            "damaged_entity": "Car" if policy_type == "Motor" else "House" if policy_type == "Property" else "Third Party",
        })

        rows["policy_coverage.csv"].append({
            "policy_coverage_identifier": policy_coverage_identifier,
            "coverage_identifier": coverage_identifier,
            "insured_entity_identifier": insured_entity_identifier,
            "coverage_status_code": "Active",
            "coverage_status_date": _iso(policy_start),
            "sum_insured": replacement_value,
            "gross_annualized_premium": premium,
            "peril_premium": round(premium * 0.28, 2),
            "annual_aggregate_limit": round(replacement_value * 1.25, 2),
            "excess": round(max(premium * 0.10, 250), 2),
            "minimum_deductible": round(max(premium * 0.05, 100), 2),
        })

        claim_count = max(
            1,
            int(sat_policy.get("Number of Active Claim", 0) or 0)
            + int(sat_policy.get("Number of Previous Claim", 0) or 0),
        )
        for claim_index in range(claim_count):
            claim_identifier = f"CLM_{batch_id}_{claim_seq:06d}"
            is_active = claim_index < int(sat_policy.get("Number of Active Claim", 0) or 0)
            policy_span_days = max((policy_end - policy_start).days, 0)
            open_offset_days = min(30 + (claim_index * 37) % 240, policy_span_days)
            open_date = _bounded_date(policy_start + timedelta(days=open_offset_days), policy_start, policy_end)
            close_date = ""
            if not is_active:
                remaining_days = max((policy_end - open_date).days, 0)
                close_offset_days = min(10 + claim_index % 20, remaining_days)
                close_date = _bounded_date(open_date + timedelta(days=close_offset_days), open_date, policy_end)
            requested = round(max(premium * (8 + claim_index), 500), 2)
            incurred = round(requested * (0.65 if is_active else 0.85), 2)
            paid = round(incurred * (0.40 if is_active else 0.92), 2)
            bodily = "Y" if policy_type == "Motor" and claim_index % 4 == 0 else "N"
            litigation = "Y" if claim_index % 9 == 0 else "N"
            fraud = "Y" if sat_policy.get("Fraud Flag") == "Y" and claim_index == 0 else "N"
            status = "Open" if is_active else "Closed"
            if litigation == "Y" and is_active:
                status = "Under Review"

            approval_date = "" if is_active else _bounded_date(open_date + timedelta(days=7), open_date, close_date)
            payment_date = "" if is_active else _bounded_date(open_date + timedelta(days=12), open_date, close_date)
            event_window_end = close_date if close_date else policy_end
            recovery_expected = round(requested * (0.18 if litigation == "Y" else 0.08), 2)
            recovery_actual = round(recovery_expected * (0.65 if not is_active else 0.20), 2)
            recovery_type = "Third Party" if litigation == "Y" else "Salvage" if policy_type == "Motor" else "Subrogation"
            risk_level = "High" if fraud == "Y" or litigation == "Y" or requested >= 15000 else "Medium" if requested >= 7000 else "Low"
            hospitalized = "Y" if bodily == "Y" and claim_index % 2 == 0 else "N"
            work_status = ["Employed", "Self Employed", "Retired", "Unemployed"][claim_seq % 4]
            if hospitalized == "Y":
                return_to_work = "Pending" if is_active else "Returned With Adjustments"
            elif bodily == "Y":
                return_to_work = "Returned Full Duty" if not is_active else "Pending"
            else:
                return_to_work = "Not Applicable"
            rejection_reason = "Fraud suspected" if fraud == "Y" and status == "Closed" else ""
            finalization_reason = "" if is_active else "Settled" if not rejection_reason else "Denied"
            reimbursable_claim = round(max(incurred - paid, 0), 2)
            reimbursable_treatment = round(requested * 0.12, 2) if bodily == "Y" else 0
            settlement_offer = round(incurred * (0.95 if not is_active else 0.75), 2)
            outstanding_reserve = round(max(incurred - paid, 0), 2)
            claims_expense = round(requested * (0.08 if litigation == "Y" else 0.04), 2)
            is_claim_suspicious = "Y" if fraud == "Y" or litigation == "Y" or risk_level == "High" else "N"
            suspected_amt = round(requested * 0.35, 2) if is_claim_suspicious == "Y" else ""
            fraud_amt = round(requested * 0.50, 2) if fraud == "Y" else ""
            is_recovery_happened = "Y" if payment_date and recovery_actual > 0 else "N"
            days_to_first_recovery = (payment_date - open_date).days if is_recovery_happened == "Y" else ""
            days_to_last_recovery = (event_window_end - open_date).days if is_recovery_happened == "Y" else ""
            litigation_duration_days = (event_window_end - open_date).days if litigation == "Y" else ""
            claim_band = "High" if requested >= 15000 else "Medium" if requested >= 7000 else "Low"
            claim_band_sort = {"Low": 1, "Medium": 2, "High": 3}[claim_band]
            claim_fraud_status = "Confirmed" if fraud == "Y" else "Under Investigation" if is_claim_suspicious == "Y" else ""
            claim_fraud_type = "Opportunistic" if fraud == "Y" else "Suspicious Pattern" if is_claim_suspicious == "Y" else ""
            claim_fraud_detection_method = "Rules Engine" if is_claim_suspicious == "Y" else ""
            recovery_band = "High" if recovery_expected >= 2500 else "Medium" if recovery_expected >= 1000 else "Low"
            recovery_category = "Legal Recovery" if litigation == "Y" else "Motor Salvage" if policy_type == "Motor" else "Property Subrogation" if policy_type == "Property" else "Liability Recovery"
            recovery_source = "Third Party" if litigation == "Y" else "Salvage Partner" if policy_type == "Motor" else "Insurer"
            indemnity_logic = "Repair network indemnity" if policy_type == "Motor" else "Cash settlement"
            emergency_service_cost = round(requested * 0.015, 2) if policy_type == "Property" and risk_level == "High" else ""
            external_fte_count = 1 + (claim_index % 2) if litigation == "Y" else ""
            rows["claim.csv"].append({
                "claim_identifier": claim_identifier,
                "claim_number": f"CLM-{batch_id}-{claim_seq:06d}",
                "third_party_claim_number": f"TP-{batch_id}-{claim_seq:06d}",
                "claim_type": policy_type,
                "claim_status": status,
                "claim_state": "In Progress" if is_active else "Finalized",
                "claim_status_date": _iso(open_date if is_active else close_date),
                "claim_open_date": _iso(open_date),
                "claim_close_date": _iso(close_date) if close_date else "",
                "claims_made_date": _iso(open_date),
                "movement_date": _iso(open_date if is_active else close_date),
                "claim_duration": (event_window_end - open_date).days,
                "claim_type_code": policy_type.upper()[:3],
                "claim_requested_amount": requested,
                "claim_specific_flag": "Y" if risk_level == "High" else "N",
                "claim_sensitivity": risk_level,
                "hospitalized_indicator": hospitalized,
                "cross_border_claim_indicator": "N",
                "intercompany_agreement_flag": "N",
                "no_claims_discount": "Y" if policy_type == "Motor" else "N",
                "outstanding_subrogation_amount": recovery_expected,
                "recovery_actual": recovery_actual,
                "recovery_type": recovery_type,
                "outstanding_reserve": outstanding_reserve,
                "claims_expense": claims_expense,
                "is_claim_suspicious": is_claim_suspicious,
                "suspected_amt": suspected_amt,
                "fraud_amt": fraud_amt,
                "is_recovery_happened": is_recovery_happened,
                "days_to_first_recovery": days_to_first_recovery,
                "days_to_last_recovery": days_to_last_recovery,
                "litigation_duration_days": litigation_duration_days,
                "claim_reason": "Collision" if policy_type == "Motor" else "Property Damage" if policy_type == "Property" else "Third Party Liability",
                "claim_channel": "Online" if claim_seq % 2 else "Call Center",
                "claim_product": policy_type,
                "claim_band": claim_band,
                "claim_band_sort": claim_band_sort,
                "claim_fraud_status": claim_fraud_status,
                "claim_fraud_type": claim_fraud_type,
                "claim_fraud_detection_method": claim_fraud_detection_method,
                "recovery_band": recovery_band,
                "recovery_category": recovery_category,
                "recovery_source": recovery_source,
                "claims_rejection_reason": rejection_reason,
                "finalization_reason": finalization_reason,
                "indemnity_logic": indemnity_logic,
                "coverage_verification_result": "Verified",
                "instruction_closure_date": _iso(close_date) if close_date else "",
                "subrogation_paid_date": _iso(payment_date) if payment_date and recovery_actual else "",
                "pre_accident_work_status": work_status,
                "return_to_work_status": return_to_work,
                "fire_brigade_fees_and_levies": emergency_service_cost,
                "claims_external_fte_count": external_fte_count,
                "total_incurred": incurred,
                "total_payment_amount": paid,
                "reimbursable_claim_amount": reimbursable_claim,
                "reimbursable_treatment_amount": reimbursable_treatment,
                "settlement_amount_pc": settlement_offer,
                "claim_amounts_description": f"Estimated {requested}; incurred {incurred}; paid {paid}",
                "policy_identifier": policy_identifier,
                "policy_coverage_identifier": policy_coverage_identifier,
                "insured_entity_identifier": insured_entity_identifier,
                "coverage_identifier": coverage_identifier,
                "claim_process_method": "Automated" if requested < 5000 else "Manual",
                "applicable_deductible_flag": "Y",
                "bodily_injury_indicator": bodily,
                "litigation_flag": litigation,
                "total_loss_flag": "Y" if requested > replacement_value * 0.6 else "N",
                "proven_claim_fraud_flag": fraud,
                "claim_approval_date": _iso(approval_date) if approval_date else "",
                "claim_payment_date": _iso(payment_date) if payment_date else "",
            })

            place = place_templates[claim_seq % len(place_templates)]
            physical_place_identifier = f"LOC_{batch_id}_{claim_seq:06d}"
            rows["physical_place.csv"].append({
                "physical_place_identifier": physical_place_identifier,
                "place_identifier": physical_place_identifier,
                "applicable_jurisdiction": applicable_jurisdiction,
                "latitude": place[1],
                "longitude": place[2],
                "sub_region": place[0],
                "surface_elevation": place[5],
                "possible_max_business_interruption": place[6],
                "point_of_interest": place[4],
                "census_zone": place[7],
                "commune_code": place[8],
                "geocoding_level": place[9],
                "municipal_district": place[10],
                "natcat_hazard_zone_scheme": place[3],
            })

            loss_date = _bounded_date(open_date - timedelta(days=claim_index % 3), policy_start, open_date)
            loss_event_identifier = f"LE_{batch_id}_{claim_seq:06d}"
            notification_lag = (open_date - loss_date).days
            is_cat = policy_type == "Property" and claim_index % 5 == 0
            rows["loss_event.csv"].append({
                "loss_event_identifier": loss_event_identifier,
                "claim_identifier": claim_identifier,
                "loss_event_name": f"{policy_type} loss event",
                "loss_event_description": f"{policy_type} claim loss reported by insured party",
                "loss_date": _iso(loss_date),
                "date_of_effective_loss": _iso(loss_date),
                "loss_event_start_date": _iso(loss_date),
                "loss_event_end_date": _iso(open_date),
                "loss_event_status": "Open" if is_active else "Closed",
                "loss_type": "Accident" if policy_type == "Motor" else "Fire" if policy_type == "Property" else "Injury",
                "loss_time": _time(loss_date),
                "loss_cause": "Collision" if policy_type == "Motor" else "Short Circuit" if policy_type == "Property" else "Negligence",
                "property_liability_loss_cause": "Third Party" if policy_type == "Liability" else "",
                "loss_category": policy_type,
                "damage_specification": "Vehicle damage" if policy_type == "Motor" else "Structural damage" if policy_type == "Property" else "Bodily injury",
                "notification_date": _iso(open_date),
                "time_to_notification": notification_lag,
                "notification_channel": "Online" if claim_seq % 2 else "Call Center",
                "days_from_event_to_fnol": notification_lag,
                "minimal_impact_flag": "Y" if requested < 5000 else "N",
                "flag_nat_cat_event": "Y" if is_cat else "N",
                "flag_catastrophe_event": "Y" if is_cat else "N",
                "fatality_indicator": "N",
                "contributory_negligence": "Y" if litigation == "Y" else "N",
                "claimant_drugs_or_alcohol": "N",
                "number_of_vehicles_involved": 2 if policy_type == "Motor" else 0,
                "number_of_people_involved": 2 if bodily == "Y" else 1,
                "number_of_injured_parties": 1 if bodily == "Y" else 0,
                "physical_place_identifier": physical_place_identifier,
            })

            if is_cat:
                rows["catastrophe.csv"].append({
                    "catastrophe_identifier": f"CAT_{batch_id}_{claim_seq:06d}",
                    "loss_event_identifier": loss_event_identifier,
                    "catastrophe_cause": "Fire",
                    "catastrophe_description": "Property loss event affecting multiple homes",
                    "catastrophe_begin_date": _iso(loss_date),
                    "catastrophe_end_date": _iso(open_date),
                    "cresta_zone": f"ZONE-{claim_seq % 9 + 1}",
                    "nat_cat_event_code": f"NAT-{claim_seq % 99:03d}",
                })

            registered_event_identifier = f"CE_{batch_id}_{event_seq:06d}"
            rows["claim_event.csv"].append({
                "claim_event_identifier": registered_event_identifier,
                "claim_identifier": claim_identifier,
                "event_type": "Claim Registered",
                "event_date": _iso(open_date),
                "event_description": "Claim reported and registered",
            })
            event_seq += 1
            last_event_identifier = registered_event_identifier
            investigation_event_identifier = registered_event_identifier
            if not is_active:
                last_event_identifier = f"CE_{batch_id}_{event_seq:06d}"
                rows["claim_event.csv"].append({
                    "claim_event_identifier": last_event_identifier,
                    "claim_identifier": claim_identifier,
                    "event_type": "Claim Settled",
                    "event_date": _iso(close_date),
                    "event_description": "Claim settlement completed",
                })
                event_seq += 1
            elif status == "Under Review":
                last_event_identifier = f"CE_{batch_id}_{event_seq:06d}"
                rows["claim_event.csv"].append({
                    "claim_event_identifier": last_event_identifier,
                    "claim_identifier": claim_identifier,
                    "event_type": "Investigation",
                    "event_date": _iso(_bounded_date(open_date + timedelta(days=2), open_date, event_window_end)),
                    "event_description": "Claim investigation ongoing",
                })
                investigation_event_identifier = last_event_identifier
                event_seq += 1

            claim_participant_identifier = f"CP_{batch_id}_{participant_seq:06d}"
            rows["claim_participant.csv"].append({
                "claim_participant_identifier": claim_participant_identifier,
                "claim_identifier": claim_identifier,
                "person_identifier": person_identifier,
                "role_in_claim": "Claimant",
                "is_primary_indicator": "Y",
                "liability_percentage": 0 if bodily != "Y" else 40,
                "injury_flag": bodily,
                "involvement_type": "Driver" if policy_type == "Motor" else "Homeowner",
                "role_start_date": _iso(open_date),
                "role_end_date": _iso(close_date) if close_date else "",
            })
            participant_seq += 1

            if not is_active:
                rows["settlement.csv"].append({
                    "settlement_identifier": f"ST_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": last_event_identifier,
                    "date_of_settlement_offer_decision": _iso(_bounded_date(open_date + timedelta(days=7), open_date, event_window_end)),
                    "date_of_final_settlement": _iso(close_date),
                    "insurers_first_settlement_offer_amount": round(paid * 0.9, 2),
                    "insurers_current_settlement_offer_amount": paid,
                    "outstanding_medical_expenses": round(max(incurred - paid, 0), 2),
                    "settlement_table": "CLAIM_SETTLEMENT",
                    "settlement_type": "Full Settlement" if paid >= incurred * 0.85 else "Partial Settlement",
                })

            condition = condition_templates[claim_seq % len(condition_templates)]
            if bodily == "Y":
                injury_identifier = f"INJ_{batch_id}_{claim_seq:06d}"
                rows["injury.csv"].append({
                    "injury_identifier": injury_identifier,
                    "claim_event_identifier": registered_event_identifier,
                    "bodily_injury_date": _iso(loss_date),
                    "injury_description": "Physical injury recorded from loss event",
                })
                rows["treatment.csv"].append({
                    "treatment_identifier": f"TR_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": registered_event_identifier,
                    "claim_participant_identifier": claim_participant_identifier,
                    "medical_condition_identifier": condition[0],
                    "treatment_code": "MED001",
                    "treatment_code_standard": "SNOMED",
                    "treatment_description": "Medical treatment for claim injury",
                    "treatment_type": "OPD" if requested < 10000 else "IPD",
                    "treatment_complexity_level": "Low" if requested < 10000 else "High",
                    "initial_treatment_date": _iso(open_date),
                    "treatment_start_date": _iso(open_date),
                    "number_of_treatments": 1 + claim_index % 3,
                    "medication_code": "RX001" if requested < 10000 else "RX010",
                    "paid_treatment_amount": round(min(paid * 0.35, requested), 2),
                    "treatment_amount_currency": "GBP",
                })
                rows["diagnosis.csv"].append({
                    "diagnosis_identifier": f"DG_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": registered_event_identifier,
                    "injury_identifier": injury_identifier,
                    "medical_condition_identifier": condition[0],
                    "diagnosis_code": "S42.3" if condition[0] == "MC_FRAC" else "S00.0",
                    "diagnosis_description": condition[1],
                    "diagnosis_priority": "Primary",
                    "diagnosis_standard": "ICD-10",
                })
                rows["medical_report.csv"].append({
                    "medical_report_identifier": f"MR_{batch_id}_{claim_seq:06d}",
                    "document_identifier": f"DOC_{batch_id}_{claim_seq:06d}",
                    "health_declaration": "Claim-related injury assessment",
                    "claim_identifier": claim_identifier,
                })
                rows["medical_assessment.csv"].append({
                    "medical_assessment_identifier": f"MA_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": registered_event_identifier,
                    "medical_report_flag": "Y",
                    "assessed_disability_degree": "0.05" if requested < 10000 else "0.20",
                    "permanent_disability_degree": "0.00" if requested < 10000 else "0.10",
                    "permanent_disability_flag": "N" if requested < 10000 else "Y",
                    "medical_report_date": _iso(_bounded_date(open_date + timedelta(days=1), open_date, event_window_end)),
                    "person_hospital_status": "Discharged" if requested < 10000 else "Hospitalized",
                    "psychological_factors": "None" if requested < 10000 else "Stress reported",
                    "medical_condition_description": condition[1],
                })
                rows["health_insurance_claim.csv"].append({
                    "health_insurance_claim_identifier": f"HCLM_{batch_id}_{claim_seq:06d}",
                    "claim_identifier": claim_identifier,
                    "date_of_first_medical_expert_report": _iso(_bounded_date(open_date + timedelta(days=1), open_date, event_window_end)),
                    "health_data_consent": "Y",
                })

            if policy_type == "Motor":
                rows["repair.csv"].append({
                    "repair_identifier": f"REP_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": last_event_identifier,
                    "claim_participant_identifier": claim_participant_identifier,
                    "number_of_spare_parts_repaired": 1 + claim_index % 5,
                    "paid_spare_parts_costs": round(paid * 0.30, 2),
                    "spare_part_repair_shop_price": round(paid * 0.33, 2),
                    "spare_part_type": "Engine Parts" if claim_index % 2 else "Body Panels",
                })

            rows["police_report.csv"].append({
                "police_report_identifier": f"PR_{batch_id}_{claim_seq:06d}",
                "claim_identifier": claim_identifier,
                "police_report_number": f"FIR-{batch_id}-{claim_seq:06d}" if policy_type == "Motor" or litigation == "Y" else "",
                "police_report_date": _iso(_bounded_date(open_date + timedelta(days=1), open_date, event_window_end)) if policy_type == "Motor" or litigation == "Y" else "",
                "police_report_flag": "Y" if policy_type == "Motor" or litigation == "Y" else "N",
            })

            if litigation == "Y":
                rows["litigation.csv"].append({
                    "litigation_identifier": f"LIT_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": last_event_identifier,
                    "claim_identifier": claim_identifier,
                    "court_matter_type": "Liability Dispute",
                    "date_of_legal_representation": _iso(_bounded_date(open_date + timedelta(days=4), open_date, event_window_end)),
                    "first_litigation_date": _iso(_bounded_date(open_date + timedelta(days=5), open_date, event_window_end)),
                    "litigation_date": _iso(_bounded_date(open_date + timedelta(days=20), open_date, event_window_end)),
                    "decision_of_court": "Pending" if is_active else "Settled",
                })

            if status == "Under Review" or fraud == "Y":
                rows["claim_investigation.csv"].append({
                    "claim_investigation_identifier": f"INV_{batch_id}_{claim_seq:06d}",
                    "claim_event_identifier": investigation_event_identifier,
                    "claim_handler_identifier": f"CH{claim_seq % 1000:04d}",
                    "claim_handler_notes": "Investigation opened due to claim risk indicators",
                    "investigator_flag": "Y",
                    "claim_investigation_start_date": _iso(_bounded_date(open_date + timedelta(days=2), open_date, event_window_end)),
                    "claim_investigation_end_date": "" if is_active else _iso(_bounded_date(open_date + timedelta(days=14), open_date, event_window_end)),
                    "fraud_indicator": fraud,
                    "claim_fraud_assessment": "Under Investigation" if is_active else "No Fraud Detected",
                })

            claim_seq += 1

    for name, schema in CLAIMS_LDM_SCHEMAS.items():
        extract_ts = _batch_pull_ts(batch_id)
        schema_rows = [
            {
                field: (
                    batch_id if field == "batch_ref"
                    else extract_ts if field == "pull_ts"
                    else "CRM" if field == "origin_sys"
                    else row.get(field, "")
                )
                for field in schema
            }
            for row in rows[name]
        ]
        write_csv(out_dir, CLAIMS_RAW_FILE_NAMES[name], schema_rows, fieldnames=schema)
    return out_dir
