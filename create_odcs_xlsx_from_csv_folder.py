from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName


DEFAULT_TEMPLATE = Path("dimdc/odcs-template-updated.xlsx")
DEFAULT_OUTPUT_DIR = Path("dimdc/outputs")
DEFAULT_RULES = Path("dimdc/rules.csv")
DEFAULT_CUSTOM_RULE_IDS = "rule_0001,rule_0007"
INPUT_MODES = {"csv", "catalog-schema", "information-schema"}

PROPERTY_COLUMNS = [
    "Property",
    "Business Name",
    "Logical Type",
    "Physical Type",
    "Example(s)",
    "Description",
    "Required",
    "Unique",
    "Classification",
    "Tags",
    "Authoritative Definition URL",
    "Authoritative Definition Type",
]

QUALITY_COLUMNS = [
    "Schema",
    "Property",
    "Quality Type",
    "Description",
    "Rule (Library)",
    "Query (SQL)",
    "Threshold Operator",
    "Threshold Value",
    "Quality Engine (Custom)",
    "Implementation (Custom)",
    "Severity",
    "Scheduler",
    "Schedule",
]


def _business_name(value: str) -> str:
    return value.replace("_", " ").replace("-", " ").title()


def _safe_sheet_title(table_name: str, used_titles: set[str]) -> str:
    raw = "Schema " + re.sub(r"[\\/*?:\[\]]", "_", table_name)
    if len(raw) <= 31 and raw not in used_titles:
        used_titles.add(raw)
        return raw
    for index in range(1, 1000):
        suffix = f"_{index}"
        candidate = raw[: 31 - len(suffix)] + suffix
        if candidate not in used_titles:
            used_titles.add(candidate)
            return candidate
    raise RuntimeError(f"Could not create a unique sheet title for {table_name}")


def _quoted_sheet_title(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _add_local_defined_name(ws: Any, name: str, cell_range: str) -> None:
    ws.defined_names.add(DefinedName(name=name, attr_text=f"{_quoted_sheet_title(ws.title)}!{cell_range}"))


def _add_schema_defined_names(ws: Any) -> None:
    ws.defined_names.clear()
    _add_local_defined_name(ws, "schema.name", "$B$5")
    _add_local_defined_name(ws, "schema.physicalType", "$B$6")
    _add_local_defined_name(ws, "schema.description", "$B$7")
    _add_local_defined_name(ws, "schema.businessName", "$B$8")
    _add_local_defined_name(ws, "schema.physicalName", "$B$9")
    _add_local_defined_name(ws, "schema.dataGranularityDescription", "$B$10")
    _add_local_defined_name(ws, "schema.tags", "$B$11")
    _add_local_defined_name(ws, "schema.properties", "$A$13:$AZ$981")


def _copy_row_style(ws: Any, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def _find_row(ws: Any, label: str) -> int:
    wanted = label.strip().lower()
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, str) and value.strip().lower() == wanted:
                return row
    raise RuntimeError(f"Could not find row label {label!r} in {ws.title!r}")


def _set_label_value(ws: Any, label: str, value: Any) -> None:
    wanted = label.strip().lower()
    for row in range(1, ws.max_row + 1):
        for col in range(1, min(ws.max_column, 6) + 1):
            cell_value = ws.cell(row, col).value
            if isinstance(cell_value, str) and cell_value.strip().lower() == wanted:
                ws.cell(row, col + 1).value = value
                return


def _clear_rows(ws: Any, start_row: int, max_col: int) -> None:
    for row in range(start_row, ws.max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None


def _read_csv_sample(path: Path, sample_rows: int) -> tuple[list[str], list[dict[str, str]], int]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        headers = [header.strip() for header in (reader.fieldnames or []) if header and header.strip()]
        rows: list[dict[str, str]] = []
        total_rows = 0
        for row in reader:
            total_rows += 1
            if len(rows) < sample_rows:
                rows.append({header: str(row.get(header, "") or "").strip() for header in headers})
    return headers, rows, total_rows


def _is_integer(values: list[str]) -> bool:
    if not values:
        return False
    for value in values:
        try:
            int(value)
        except ValueError:
            return False
    return True


def _is_number(values: list[str]) -> bool:
    if not values:
        return False
    for value in values:
        try:
            float(value)
        except ValueError:
            return False
    return True


def _is_boolean(values: list[str]) -> bool:
    valid = {"true", "false", "t", "f", "yes", "no", "y", "n", "0", "1"}
    return bool(values) and all(value.lower() in valid for value in values)


def _is_datetime(values: list[str]) -> bool:
    if not values:
        return False
    for value in values:
        text = value.replace("Z", "+00:00")
        try:
            datetime.fromisoformat(text)
        except ValueError:
            return False
    return True


def _is_date(values: list[str]) -> bool:
    if not values:
        return False
    for value in values:
        try:
            datetime.strptime(value[:10], "%Y-%m-%d")
        except ValueError:
            return False
    return True


def _infer_types(values: list[str]) -> tuple[str, str]:
    non_empty = [value for value in values if value != ""]
    if _is_boolean(non_empty):
        return "boolean", "boolean"
    if _is_integer(non_empty):
        return "integer", "bigint"
    if _is_number(non_empty):
        return "number", "decimal(18,2)"
    if _is_datetime(non_empty):
        return "timestamp", "timestamp"
    if _is_date(non_empty):
        return "date", "date"

    max_len = max((len(value) for value in non_empty), default=64)
    size = 32 if max_len <= 32 else 64 if max_len <= 64 else 128 if max_len <= 128 else 512
    return "string", f"varchar({size})"


def _classification(column: str) -> str:
    normalized = column.lower()
    sensitive_tokens = [
        "email",
        "phone",
        "mobile",
        "name",
        "address",
        "dob",
        "birth",
        "gender",
        "nationality",
        "postcode",
        "zipcode",
        "zip",
    ]
    if any(token in normalized for token in sensitive_tokens):
        return "confidential"
    if normalized.endswith("_sk") or normalized == "sk":
        return "restricted"
    return "internal"


def _logical_type_from_database_type(data_type: str) -> tuple[str, str]:
    normalized = data_type.lower().strip()
    if any(token in normalized for token in ("tinyint", "smallint", "int", "bigint")):
        return "integer", data_type
    if any(token in normalized for token in ("decimal", "numeric", "double", "float", "real")):
        return "number", data_type
    if "bool" in normalized:
        return "boolean", data_type
    if "timestamp" in normalized or "datetime" in normalized:
        return "timestamp", data_type
    if normalized == "date" or normalized.startswith("date"):
        return "date", data_type
    if "array" in normalized:
        return "array", data_type
    if "struct" in normalized or "map" in normalized:
        return "object", data_type
    return "string", data_type or "string"


def _infer_primary_key(table_name: str, columns: list[str], rows: list[dict[str, str]]) -> str | None:
    candidates = [
        f"{table_name}_id",
        f"{table_name}_identifier",
        f"{table_name}_sk",
        "id",
        "identifier",
    ]
    lower_to_actual = {column.lower(): column for column in columns}
    for candidate in candidates:
        if candidate in lower_to_actual:
            return lower_to_actual[candidate]

    for suffix in ("_identifier", "_id", "_sk"):
        for column in columns:
            if column.lower().endswith(suffix):
                return column

    for column in columns:
        values = [row.get(column, "") for row in rows if row.get(column, "")]
        if values and len(values) == len(set(values)):
            return column
    return None


def _column_required(column: str, pk: str | None, rows: list[dict[str, str]]) -> bool:
    if column == pk:
        return True
    if column in {"batch_ref", "pull_ts", "origin_sys"}:
        return True
    if column.lower().endswith(("_id", "_identifier", "_sk")):
        return True
    values = [row.get(column, "") for row in rows]
    return bool(values) and all(value != "" for value in values)


def _table_metadata(path: Path, sample_rows: int) -> dict[str, Any]:
    columns, rows, row_count = _read_csv_sample(path, sample_rows)
    table_name = path.stem
    pk = _infer_primary_key(table_name, columns, rows)
    properties = []
    for column in columns:
        values = [row.get(column, "") for row in rows]
        logical_type, physical_type = _infer_types(values)
        properties.append({
            "Property": column,
            "Business Name": _business_name(column),
            "Logical Type": logical_type,
            "Physical Type": physical_type,
            "Example(s)": next((value for value in values if value), ""),
            "Description": f"Source field {column}.",
            "Required": "TRUE" if _column_required(column, pk, rows) else "FALSE",
            "Unique": "TRUE" if column == pk else "FALSE",
            "Classification": _classification(column),
            "Tags": "primary_key" if column == pk else "",
            "Authoritative Definition URL": "",
            "Authoritative Definition Type": "",
        })
    return {
        "table_name": table_name,
        "physical_name": path.name,
        "path": path,
        "columns": columns,
        "rows": rows,
        "row_count": row_count,
        "primary_key": pk,
        "properties": properties,
    }


def _property_rows_from_information_schema(
    table_name: str,
    columns: list[dict[str, Any]],
    pk: str | None,
) -> list[dict[str, Any]]:
    properties = []
    for column in columns:
        column_name = str(column["column_name"])
        data_type = str(column.get("data_type") or "string")
        logical_type, physical_type = _logical_type_from_database_type(data_type)
        is_pk = column_name == pk
        properties.append({
            "Property": column_name,
            "Business Name": _business_name(column_name),
            "Logical Type": logical_type,
            "Physical Type": physical_type,
            "Example(s)": "",
            "Description": str(column.get("comment") or f"Information schema field {column_name}."),
            "Required": "TRUE" if is_pk or column_name.lower().endswith(("_id", "_identifier", "_sk")) else "FALSE",
            "Unique": "TRUE" if is_pk else "FALSE",
            "Classification": _classification(column_name),
            "Tags": "primary_key" if is_pk else "",
            "Authoritative Definition URL": "",
            "Authoritative Definition Type": "",
        })
    return properties


def _table_metadata_from_information_schema(
    table_name: str,
    columns: list[dict[str, Any]],
    table_comment: str = "",
) -> dict[str, Any]:
    column_names = [str(column["column_name"]) for column in columns]
    pk = _infer_primary_key(table_name, column_names, [])
    return {
        "table_name": table_name,
        "physical_name": table_name,
        "path": None,
        "columns": column_names,
        "rows": [],
        "row_count": None,
        "primary_key": pk,
        "description": table_comment or f"Table {table_name} from information_schema.",
        "properties": _property_rows_from_information_schema(table_name, columns, pk),
    }


def _quote_identifier(identifier: str) -> str:
    return "`" + identifier.replace("`", "``") + "`"


def _parse_catalog_schema(value: str) -> tuple[str, str]:
    parts = [part.strip() for part in str(value or "").split(".") if part.strip()]
    if len(parts) != 2:
        raise RuntimeError("catalog-schema mode expects --input-path in the form <catalog>.<schema>.")
    return parts[0], parts[1]


def _fetch_information_schema_tables(
    catalog_name: str,
    schema_name: str,
    server_hostname: str | None,
    http_path: str | None,
    access_token: str | None,
    table_filter: list[str],
) -> list[dict[str, Any]]:
    try:
        from databricks import sql
    except ImportError as exc:
        raise RuntimeError(
            "information-schema mode requires databricks-sql-connector. "
            "Install it with: pip install databricks-sql-connector"
        ) from exc

    missing = []
    if not server_hostname:
        missing.append("DATABRICKS_SERVER_HOSTNAME or --databricks-server-hostname")
    if not http_path:
        missing.append("DATABRICKS_HTTP_PATH or --databricks-http-path")
    if not access_token:
        missing.append("DATABRICKS_TOKEN or --databricks-token")
    if missing:
        raise RuntimeError("Missing Databricks connection settings: " + ", ".join(missing))

    catalog_ref = _quote_identifier(catalog_name)
    table_filter_set = {table.lower() for table in table_filter}
    columns_query = f"""
        SELECT table_name, column_name, data_type, ordinal_position, comment
        FROM {catalog_ref}.information_schema.columns
        WHERE table_schema = ?
        ORDER BY table_name, ordinal_position
    """
    tables_query = f"""
        SELECT table_name, comment
        FROM {catalog_ref}.information_schema.tables
        WHERE table_schema = ?
    """

    with sql.connect(
        server_hostname=server_hostname,
        http_path=http_path,
        access_token=access_token,
    ) as connection:
        with connection.cursor() as cursor:
            cursor.execute(tables_query, (schema_name,))
            table_comments = {
                str(row[0]): str(row[1] or "")
                for row in cursor.fetchall()
                if not table_filter_set or str(row[0]).lower() in table_filter_set
            }
            cursor.execute(columns_query, (schema_name,))
            columns_by_table: dict[str, list[dict[str, Any]]] = {}
            for row in cursor.fetchall():
                table_name = str(row[0])
                if table_filter_set and table_name.lower() not in table_filter_set:
                    continue
                columns_by_table.setdefault(table_name, []).append({
                    "table_name": table_name,
                    "column_name": str(row[1]),
                    "data_type": str(row[2] or "string"),
                    "ordinal_position": row[3],
                    "comment": str(row[4] or ""),
                })

    return [
        _table_metadata_from_information_schema(table_name, columns, table_comments.get(table_name, ""))
        for table_name, columns in sorted(columns_by_table.items())
    ]


def _write_fundamentals(wb: Any, input_path: Path, contract_name: str, domain: str, data_product: str) -> None:
    if "Fundamentals" not in wb.sheetnames:
        return
    ws = wb["Fundamentals"]
    values = {
        "Kind": "DataContract",
        "API Version": "v3.1.0",
        "ID": re.sub(r"[^a-zA-Z0-9_-]+", "-", contract_name.lower()).strip("-"),
        "Name": contract_name,
        "Version": "1.0.0",
        "Status": "draft",
        "Domain": domain,
        "Data Product": data_product,
        "Tenant": "",
        "Purpose": f"Ad hoc ODCS Excel generated from {input_path}.",
        "Limitations": "Generated from CSV headers and sampled values; review inferred types and key columns before approval.",
        "Usage": "Used to bootstrap an ODCS data contract workbook from a folder of CSV files.",
    }
    for label, value in values.items():
        _set_label_value(ws, label, value)


def _write_schema_sheet(ws: Any, table: dict[str, Any], tags: str) -> None:
    _add_schema_defined_names(ws)
    table_name = table["table_name"]
    _set_label_value(ws, "Name", table_name)
    _set_label_value(ws, "Type", "table")
    _set_label_value(ws, "Description", table.get("description") or f"CSV source records for {table_name}.")
    _set_label_value(ws, "Business Name", _business_name(table_name))
    _set_label_value(ws, "Physical Name", table["physical_name"])
    grain = f"One row per {table['primary_key']}." if table["primary_key"] else "One row per CSV record."
    _set_label_value(ws, "Data Granularity", grain)
    _set_label_value(ws, "Tags", tags)

    header_row = _find_row(ws, "Property")
    data_start = header_row + 1
    _clear_rows(ws, data_start, max(len(PROPERTY_COLUMNS), ws.max_column))
    for row_index, prop in enumerate(table["properties"], start=data_start):
        _copy_row_style(ws, data_start, row_index, len(PROPERTY_COLUMNS))
        for col_index, key in enumerate(PROPERTY_COLUMNS, start=1):
            ws.cell(row_index, col_index).value = prop.get(key)


def _load_rule_catalog(rules_path: Path | None) -> dict[str, dict[str, str]]:
    if not rules_path or not rules_path.exists():
        return {}
    rules = pd.read_csv(rules_path).fillna("")
    return {
        str(row["rule_id"]).strip(): {key: str(value).strip() for key, value in row.items()}
        for row in rules.to_dict("records")
        if str(row.get("rule_id", "")).strip()
    }


def _custom_rule_quality_row(table: dict[str, Any], rule: dict[str, str]) -> dict[str, Any] | None:
    rule_id = rule.get("rule_id", "")
    rule_name = rule.get("rule_name", rule_id)
    if not table["primary_key"]:
        return None
    return {
        "Schema": table["table_name"],
        "Property": table["primary_key"],
        "Quality Type": "custom",
        "Description": f"{rule_id}: {rule_name} - {rule.get('query_desc', '')}",
        "Rule (Library)": rule_id,
        "Query (SQL)": "",
        "Threshold Operator": "",
        "Threshold Value": "",
        "Quality Engine (Custom)": "business_rule_catalog",
        "Implementation (Custom)": f"{rule_id}: {rule_name}",
        "Severity": rule.get("rule_type") or "Warning",
        "Scheduler": "",
        "Schedule": "",
    }


def _write_quality(
    wb: Any,
    tables: list[dict[str, Any]],
    rule_catalog: dict[str, dict[str, str]],
    custom_rule_ids: list[str],
) -> None:
    if "Quality" not in wb.sheetnames:
        return
    ws = wb["Quality"]
    header_row = _find_row(ws, "Schema")
    data_start = header_row + 1
    _clear_rows(ws, data_start, len(QUALITY_COLUMNS))

    row_number = data_start
    quality_rows: list[dict[str, Any]] = []
    for table in tables:
        quality_rows.append({
            "Schema": table["table_name"],
            "Property": "",
            "Quality Type": "custom",
            "Description": "Table must contain at least one row.",
            "Rule (Library)": "rowCount",
            "Query (SQL)": "",
            "Threshold Operator": "",
            "Threshold Value": "",
            "Quality Engine (Custom)": "odcs_metric",
            "Implementation (Custom)": "rowCount mustBeGreaterThan 0",
            "Severity": "Warning",
            "Scheduler": "",
            "Schedule": "",
        })
        if table["primary_key"]:
            quality_rows.append({
                "Schema": table["table_name"],
                "Property": table["primary_key"],
                "Quality Type": "custom",
                "Description": "Primary key must not contain null values.",
                "Rule (Library)": "nullValues",
                "Query (SQL)": "",
                "Threshold Operator": "",
                "Threshold Value": "",
                "Quality Engine (Custom)": "odcs_metric",
                "Implementation (Custom)": "nullValues mustBe 0",
                "Severity": "Warning",
                "Scheduler": "",
                "Schedule": "",
            })
            quality_rows.append({
                "Schema": table["table_name"],
                "Property": table["primary_key"],
                "Quality Type": "custom",
                "Description": "Primary key must not contain duplicate values.",
                "Rule (Library)": "duplicateValues",
                "Query (SQL)": "",
                "Threshold Operator": "",
                "Threshold Value": "",
                "Quality Engine (Custom)": "odcs_metric",
                "Implementation (Custom)": "duplicateValues mustBe 0",
                "Severity": "Warning",
                "Scheduler": "",
                "Schedule": "",
            })
            for rule_id in custom_rule_ids:
                rule = rule_catalog.get(rule_id)
                if not rule:
                    continue
                row = _custom_rule_quality_row(table, rule)
                if row:
                    quality_rows.append(row)

    for row in quality_rows:
        _copy_row_style(ws, data_start, row_number, len(QUALITY_COLUMNS))
        for col_index, key in enumerate(QUALITY_COLUMNS, start=1):
            ws.cell(row_number, col_index).value = row.get(key)
        row_number += 1


def create_odcs_xlsx_from_csv_folder(
    input_path: Path,
    output_path: Path,
    template_path: Path,
    sample_rows: int,
    contract_name: str,
    domain: str,
    data_product: str,
    tags: str,
    rules_path: Path | None,
    custom_rule_ids: list[str],
    input_mode: str = "csv",
    catalog_name: str | None = None,
    schema_name: str | None = None,
    databricks_server_hostname: str | None = None,
    databricks_http_path: str | None = None,
    databricks_token: str | None = None,
    table_filter: list[str] | None = None,
) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"ODCS Excel template not found: {template_path}")

    if input_mode not in INPUT_MODES:
        raise RuntimeError(f"Unsupported input mode {input_mode!r}. Use one of: {', '.join(sorted(INPUT_MODES))}")

    if input_mode in {"catalog-schema", "information-schema"}:
        if (not catalog_name or not schema_name) and str(input_path):
            parsed_catalog, parsed_schema = _parse_catalog_schema(str(input_path))
            catalog_name = catalog_name or parsed_catalog
            schema_name = schema_name or parsed_schema
        if not catalog_name or not schema_name:
            raise RuntimeError("catalog-schema mode requires --input-path <catalog>.<schema> or --catalog-name and --schema-name.")
        tables = _fetch_information_schema_tables(
            catalog_name=catalog_name,
            schema_name=schema_name,
            server_hostname=databricks_server_hostname or os.getenv("DATABRICKS_SERVER_HOSTNAME"),
            http_path=databricks_http_path or os.getenv("DATABRICKS_HTTP_PATH"),
            access_token=databricks_token or os.getenv("DATABRICKS_TOKEN"),
            table_filter=table_filter or [],
        )
        if not tables:
            raise RuntimeError(f"No tables found in {catalog_name}.{schema_name}")
        source_label = Path(f"{catalog_name}.{schema_name}")
    else:
        if not input_path.exists():
            raise FileNotFoundError(f"Input path does not exist: {input_path}")
        if input_path.is_file():
            if input_path.suffix.lower() != ".csv":
                raise RuntimeError(f"Input file must be a CSV file: {input_path}")
            csv_files = [input_path]
            source_label = input_path.parent
        else:
            csv_files = sorted(input_path.glob("*.csv"))
            source_label = input_path

        if not csv_files:
            raise RuntimeError(f"No CSV files found in {input_path}")
        tables = [_table_metadata(path, sample_rows) for path in csv_files]

    rule_catalog = _load_rule_catalog(rules_path)
    wb = load_workbook(template_path)
    _write_fundamentals(wb, source_label, contract_name, domain, data_product)

    schema_templates = [name for name in wb.sheetnames if name.startswith("Schema")]
    if not schema_templates:
        raise RuntimeError("Template does not contain any Schema sheet.")
    template_ws = wb[schema_templates[0]]
    used_titles = {name for name in wb.sheetnames if not name.startswith("Schema")}

    created_sheets = []
    for table in tables:
        ws = wb.copy_worksheet(template_ws)
        ws.title = _safe_sheet_title(table["table_name"], used_titles)
        _write_schema_sheet(ws, table, tags)
        created_sheets.append(ws.title)

    for name in schema_templates:
        if name in wb.sheetnames:
            del wb[name]

    fundamentals_index = wb.sheetnames.index("Fundamentals") + 1 if "Fundamentals" in wb.sheetnames else 0
    for offset, name in enumerate(created_sheets):
        current_index = wb.sheetnames.index(name)
        wb.move_sheet(wb[name], offset=-(current_index - fundamentals_index - offset))

    _write_quality(wb, tables, rule_catalog, custom_rule_ids)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _default_output(input_path: Path) -> Path:
    safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", input_path.stem if input_path.is_file() else input_path.name).strip("_") or "csv_input"
    return DEFAULT_OUTPUT_DIR / f"{safe_name}_odcs_contract.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser(description="Create an ODCS Excel workbook from CSV files or Databricks catalog.schema.")
    parser.add_argument(
        "--input-mode",
        choices=sorted(INPUT_MODES),
        default="csv",
        help="csv reads a CSV file/folder. catalog-schema reads table/column metadata from Databricks information_schema for <catalog>.<schema>. information-schema is a backwards-compatible alias.",
    )
    parser.add_argument("--input-path", help="CSV file/folder for csv mode, or <catalog>.<schema> for catalog-schema mode.")
    parser.add_argument("--output", help="Output XLSX path. Defaults to dimdc/outputs/<folder>_odcs_contract.xlsx.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="ODCS Excel template path.")
    parser.add_argument("--sample-rows", type=int, default=1000, help="Rows sampled per CSV for type/key inference.")
    parser.add_argument("--contract-name", default="Ad Hoc CSV Folder ODCS Data Contract")
    parser.add_argument("--domain", default="adhoc")
    parser.add_argument("--data-product", default="csv_folder")
    parser.add_argument("--tags", default="adhoc, csv, source_feed")
    parser.add_argument("--catalog-name", help="Optional Databricks catalog name for catalog-schema mode.")
    parser.add_argument("--schema-name", help="Optional Databricks schema/database name for catalog-schema mode.")
    parser.add_argument("--table-names", default="", help="Optional comma-separated table list for catalog-schema mode.")
    parser.add_argument("--databricks-server-hostname", help="Databricks server hostname. Defaults to DATABRICKS_SERVER_HOSTNAME.")
    parser.add_argument("--databricks-http-path", help="Databricks SQL warehouse HTTP path. Defaults to DATABRICKS_HTTP_PATH.")
    parser.add_argument("--databricks-token", help="Databricks token. Defaults to DATABRICKS_TOKEN.")
    parser.add_argument("--rules", default=str(DEFAULT_RULES), help="Rule catalog CSV for custom quality rule references.")
    parser.add_argument(
        "--custom-rule-ids",
        default=DEFAULT_CUSTOM_RULE_IDS,
        help="Comma-separated rule IDs from --rules to add for inferred primary keys. Use empty string to disable.",
    )
    args = parser.parse_args()

    if args.input_mode == "csv" and not args.input_path:
        parser.error("--input-path is required when --input-mode csv")
    if args.input_mode in {"catalog-schema", "information-schema"} and not args.input_path and (not args.catalog_name or not args.schema_name):
        parser.error("--input-path <catalog>.<schema> or both --catalog-name and --schema-name are required for catalog-schema mode")

    input_path = Path(args.input_path) if args.input_path else Path(f"{args.catalog_name}.{args.schema_name}")
    output_path = Path(args.output) if args.output else _default_output(input_path)
    custom_rule_ids = [rule_id.strip() for rule_id in args.custom_rule_ids.split(",") if rule_id.strip()]
    table_filter = [table.strip() for table in args.table_names.split(",") if table.strip()]
    try:
        output = create_odcs_xlsx_from_csv_folder(
            input_path=input_path,
            output_path=output_path,
            template_path=Path(args.template),
            sample_rows=args.sample_rows,
            contract_name=args.contract_name,
            domain=args.domain,
            data_product=args.data_product,
            tags=args.tags,
            rules_path=Path(args.rules) if args.rules else None,
            custom_rule_ids=custom_rule_ids,
            input_mode=args.input_mode,
            catalog_name=args.catalog_name,
            schema_name=args.schema_name,
            databricks_server_hostname=args.databricks_server_hostname,
            databricks_http_path=args.databricks_http_path,
            databricks_token=args.databricks_token,
            table_filter=table_filter,
        )
    except (FileNotFoundError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1) from None
    print(f"Generated: {output}")


if __name__ == "__main__":
    main()
