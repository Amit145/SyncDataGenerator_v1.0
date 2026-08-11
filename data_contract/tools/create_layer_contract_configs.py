"""Create Bronze/Silver/Gold data contract config YAML from layer metadata workbooks."""

from __future__ import annotations

import json
import copy
import re
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook


INPUT_DIR = Path("data_contract/inputs")
OUTPUT_DIR = Path("data_contract/tools")

RULES_FILE = INPUT_DIR / "dq_rules.xlsx"
BRONZE_TABLES_FILE = INPUT_DIR / "bronze_tables_c360.xlsx"
SILVER_TABLES_FILE = INPUT_DIR / "silver_tables_c360.xlsx"
GOLD_TABLES_FILE = INPUT_DIR / "gold_tables_c360.xlsx"


def clean_text(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.replace("_x000D_", "\n").strip()
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n[ \t]+", "\n", text)
        return text if text.lower() not in {"", "null", "none"} else None
    return value


def read_sheet(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [clean_text(worksheet.cell(1, col).value) for col in range(1, worksheet.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for row_num in range(2, worksheet.max_row + 1):
        row = {
            str(headers[col - 1]): clean_text(worksheet.cell(row_num, col).value)
            for col in range(1, worksheet.max_column + 1)
            if headers[col - 1]
        }
        if any(value is not None for value in row.values()):
            rows.append(row)
    return rows


def key_list(value: Any) -> list[str]:
    value = clean_text(value)
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    separator = "|" if "|" in text else ","
    return [part.strip() for part in text.split(separator) if part.strip()]


def key_value(value: Any) -> str | list[str] | None:
    keys = key_list(value)
    if not keys:
        return None
    return keys[0] if len(keys) == 1 else keys


def title_name(value: str) -> str:
    return " ".join(part.capitalize() for part in value.replace("*", "").replace(".csv", "").split("_"))


def parse_expected(value: Any) -> Any:
    value = clean_text(value)
    if value is None:
        return None
    if isinstance(value, dict):
        return value
    try:
        return json.loads(str(value))
    except json.JSONDecodeError:
        return yaml.safe_load(str(value))


def load_rules() -> dict[str, dict[str, Any]]:
    rules = {}
    for row in read_sheet(RULES_FILE):
        rule_id = row.get("rule_id")
        if not rule_id:
            continue
        rules[str(rule_id)] = {
            "ruleId": str(rule_id),
            "ruleName": row.get("rule_name") or str(rule_id),
            "type": "sql",
            "query": row.get("query_text") or "",
            "mustBe": parse_expected(row.get("expected_outcome")),
            "dimension": row.get("rule_category") or "validity",
            "severity": str(row.get("rule_type") or "warning").lower(),
            "description": row.get("query_desc") or row.get("rule_name") or str(rule_id),
            "failureAction": "warn",
        }
    return rules


def referenced_columns_for_rule(rule_id: str, pk: list[str], bk: list[str]) -> list[str]:
    if rule_id == "rule_0003":
        return pk
    if rule_id in {"rule_0001", "rule_0007", "rule_0002", "rule_0005", "rule_0006"}:
        return bk
    if rule_id == "rule_0004":
        return ["effective_from_ts", "effective_to_ts"]
    return []


def rule_from_catalog(rule_catalog: dict[str, dict[str, Any]], rule_id: str, pk: list[str], bk: list[str]) -> dict[str, Any]:
    rule = copy.deepcopy(rule_catalog[rule_id])
    refs = referenced_columns_for_rule(rule_id, pk, bk)
    if refs:
        rule["referencedColumns"] = list(refs)
    return rule


def required_columns(*groups: list[str]) -> dict[str, dict[str, bool]]:
    columns: dict[str, dict[str, bool]] = {}
    for group in groups:
        for column in group:
            columns.setdefault(column, {"required": True})
    return columns


def base_config(layer: str) -> dict[str, Any]:
    layer_title = layer.capitalize()
    return {
        "purpose": f"Validate C360 {layer} layer tables before downstream processing.",
        "limitations": "Generated from layer metadata workbooks; review business descriptions and column-level rules when physical schema is available.",
        "usage": f"Used as a {layer} layer contract config for ODCS contract generation and validation.",
        "frequency": 1,
        "frequencyUnit": "d",
        "latency": 24,
        "latencyUnit": "h",
        "retention": 7,
        "retentionUnit": "y",
        "teamName": "Data Modelling & Engineering CoE",
        "firstLevelApprovers": f"C360 {layer_title} Data Owner",
        "secondLevelApprovers": "Data Governance",
        "supportChannel": "coe-data-contracts",
        "supportUrl": "mailto:coe-data-contracts@example-internal",
        "tables": {},
    }


def build_bronze_config(rule_catalog: dict[str, dict[str, Any]]) -> dict[str, Any]:
    config = base_config("bronze")
    for row in read_sheet(BRONZE_TABLES_FILE):
        if row.get("is_active") is not True:
            continue
        table_name = str(row["target_table"]).strip()
        pk = key_list(row.get("primary_keys"))
        bk = key_list(row.get("business_key")) or pk
        source_object = str(row.get("source_object_name") or "").replace("*", "")
        config["tables"][table_name] = {
            "name": table_name,
            "primaryKey": key_value(pk),
            "businessKey": key_value(bk),
            "businessName": title_name(table_name),
            "grain": f"One row per {', '.join(bk or pk)} per batch.",
            "columns": required_columns(pk, bk),
            "metadata": {
                "bronzeObjectId": row.get("bronze_object_id"),
                "sourceId": row.get("source_id"),
                "sourceObjectName": source_object,
                "landingPath": row.get("landing_path"),
                "fileFormat": row.get("file_format"),
                "writeMode": row.get("write_mode"),
                "schemaEvolution": row.get("schema_evolution"),
            },
        }
    return config


def build_silver_or_gold_config(layer: str, table_file: Path, rule_catalog: dict[str, dict[str, Any]]) -> dict[str, Any]:
    config = base_config(layer)
    for row in read_sheet(table_file):
        if row.get("is_active") is not True:
            continue
        table_name = str(row["table_name"]).strip()
        pk = key_list(row.get("table_pk"))
        bk = key_list(row.get("table_bk")) or pk
        source_table = row.get("source_table")
        source_pk = key_list(row.get("source_pk"))
        table_rules = []

        if row.get("scd2") is True:
            for rule_id in ("rule_0004", "rule_0005", "rule_0006"):
                table_rules.append(rule_from_catalog(rule_catalog, rule_id, pk, bk))

        if layer == "gold" and source_table and source_pk:
            table_rules.append(rule_from_catalog(rule_catalog, "rule_0002", pk, bk))

        table_cfg: dict[str, Any] = {
            "name": table_name,
            "primaryKey": key_value(pk),
            "businessKey": key_value(bk),
            "businessName": title_name(table_name),
            "grain": f"One row per {', '.join(bk or pk)}.",
            "columns": required_columns(pk, bk, ["effective_from_ts", "effective_to_ts"] if row.get("scd2") is True else []),
            "metadata": {
                "tableId": row.get("table_id"),
                "catalogName": row.get("catalog_name"),
                "schemaName": row.get("schema_name"),
                "stage": row.get("stage") or layer,
                "scd2": row.get("scd2"),
            },
        }
        if source_table:
            table_cfg["sourceTable"] = source_table
        if source_pk:
            table_cfg["sourcePrimaryKey"] = key_value(source_pk)
        if table_rules:
            table_cfg["tableRules"] = table_rules
        config["tables"][table_name] = table_cfg
    return config


def write_yaml(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(payload, sort_keys=False, allow_unicode=False), encoding="utf-8")
    print(f"Wrote {path}")


def main() -> None:
    rule_catalog = load_rules()
    outputs = {
        OUTPUT_DIR / "bronze_c360_contract_config.yaml": build_bronze_config(rule_catalog),
        OUTPUT_DIR / "silver_c360_contract_config.yaml": build_silver_or_gold_config("silver", SILVER_TABLES_FILE, rule_catalog),
        OUTPUT_DIR / "gold_c360_contract_config.yaml": build_silver_or_gold_config("gold", GOLD_TABLES_FILE, rule_catalog),
    }
    for path, payload in outputs.items():
        write_yaml(path, payload)


if __name__ == "__main__":
    main()
