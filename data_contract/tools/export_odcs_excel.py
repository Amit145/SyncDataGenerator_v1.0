"""Export an ODCS YAML contract into the official ODCS Excel template shape."""

from __future__ import annotations

import json
import re
from copy import copy
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName


DEFAULT_EXCEL_TEMPLATE = Path("dc_nb/odcs-template.xlsx")
RULE_NAMES = {
    "rule_0007": "BK Not Null Check",
    "rule_0003": "Not Null Check",
    "rule_0001": "BK Check",
    "rule_0002": "Missing Keys Check",
}


def _scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _join_values(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return str(value)


def _first_numeric_value(value: Any) -> int | float:
    if isinstance(value, dict):
        for item in value.values():
            if isinstance(item, (int, float)) and not isinstance(item, bool):
                return item
    if isinstance(value, list):
        for item in value:
            if isinstance(item, (int, float)) and not isinstance(item, bool):
                return item
    return 0


def _sql_literal(value: Any) -> str:
    text = str(value).replace("'", "''")
    return f"'{text}'"


def _library_quality_query(metric: str | None, prop_name: str | None, quality: dict[str, Any]) -> str:
    if metric == "rowCount":
        return "SELECT COUNT(*) AS row_count FROM {full_table_name}"
    if prop_name and metric in {"nullValues", "missingValues"}:
        return f"SELECT COUNT(*) AS null_count FROM {{full_table_name}} WHERE {prop_name} IS NULL"
    if prop_name and metric == "duplicateValues":
        return f"SELECT COUNT(*) - COUNT(DISTINCT {prop_name}) AS duplicate_count FROM {{full_table_name}}"
    if prop_name and metric == "invalidValues":
        arguments = quality.get("arguments") or {}
        valid_values = arguments.get("validValues") if isinstance(arguments, dict) else None
        if valid_values:
            values = ", ".join(_sql_literal(value) for value in valid_values)
            return (
                f"SELECT COUNT(*) AS invalid_count FROM {{full_table_name}} "
                f"WHERE {prop_name} IS NOT NULL AND {prop_name} NOT IN ({values})"
            )
        return f"SELECT COUNT(*) AS invalid_count FROM {{full_table_name}} WHERE {prop_name} IS NULL"
    return "SELECT COUNT(*) AS check_count FROM {full_table_name}"


def _quality_operator(quality: dict[str, Any]) -> tuple[str | None, Any]:
    operator_keys = [
        "mustBe",
        "mustNotBe",
        "mustBeGreaterThan",
        "mustBeGreaterOrEqualTo",
        "mustBeLessThan",
        "mustBeLessOrEqualTo",
        "mustBeBetween",
    ]
    operator = next((key for key in operator_keys if key in quality), None)
    return operator, quality.get(operator) if operator else None


def _rule_no(quality: dict[str, Any]) -> str | None:
    for key in ("id", "name", "description"):
        value = quality.get(key)
        if not value:
            continue
        match = re.search(r"\b(rule[_-]?\d+|dq_[A-Za-z0-9_]+)\b", str(value), flags=re.IGNORECASE)
        if match:
            return match.group(1)
    return quality.get("id") or quality.get("name") or quality.get("metric")


def _referenced_columns(quality: dict[str, Any], prop_name: str | None) -> str | None:
    if prop_name:
        return prop_name
    for item in quality.get("customProperties") or []:
        if item.get("property") == "referencedColumns":
            return _join_values(item.get("value"))
    return None


def _query_reference(schema_name: str, prop_name: str | None, quality: dict[str, Any]) -> str:
    rule = _rule_no(quality) or "n/a"
    columns = _referenced_columns(quality, prop_name)
    if columns:
        return f"rule={rule}; column={columns}"
    return f"rule={rule}; table={schema_name}"


def _quality_type(quality: dict[str, Any]) -> str:
    return quality.get("type") or ("library" if quality.get("metric") else "sql" if quality.get("query") else "custom")


def _rule_label(quality: dict[str, Any]) -> str | None:
    rule = _rule_no(quality)
    name = quality.get("name") or RULE_NAMES.get(str(rule or "").lower())
    if not name and quality.get("description"):
        match = re.search(r"\b(rule[_-]?\d+|dq_[A-Za-z0-9_]+)\s*:\s*([^-]+)", str(quality["description"]))
        if match:
            name = match.group(2).strip()
    if name and rule and str(name) != str(rule):
        return f"{rule} - {name}"
    if rule:
        return str(rule)
    return name


def _copy_row_style(ws: Any, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def _sheet_title(table_name: str, used_titles: set[str]) -> str:
    raw = "Schema " + re.sub(r"[\\/*?:\[\]]", "_", table_name)
    candidates = [raw, raw[:31]]
    for candidate in candidates:
        if candidate and len(candidate) <= 31 and candidate not in used_titles:
            used_titles.add(candidate)
            return candidate

    base = raw[:31]
    for suffix_number in range(1, 100):
        suffix = f"_{suffix_number}"
        candidate = base[: 31 - len(suffix)] + suffix
        if candidate not in used_titles:
            used_titles.add(candidate)
            return candidate

    raise RuntimeError(f"Could not create a unique Excel sheet title for {table_name}")


def _quoted_sheet_title(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _add_local_defined_name(ws: Any, name: str, cell_range: str) -> None:
    ws.defined_names.add(DefinedName(name=name, attr_text=f"{_quoted_sheet_title(ws.title)}!{cell_range}"))


def _add_schema_defined_names(ws: Any) -> None:
    ws.defined_names.clear()
    _add_local_defined_name(ws, "schema.name", "$B$5")
    _add_local_defined_name(ws, "schema.physicalType", "$B$6")
    _add_local_defined_name(ws, "schema.description", "$B$7")
    _add_local_defined_name(ws, "schema.businessName", "$B$8")
    _add_local_defined_name(ws, "schema.physicalName", "$B$9")
    _add_local_defined_name(ws, "schema.dataGranularityDescription", "$B$10")
    _add_local_defined_name(ws, "schema.tags", "$B$11")
    _add_local_defined_name(ws, "schema.properties", "$A$13:$AZ$981")


def _clear_rows(ws: Any, start_row: int, max_col: int) -> None:
    for row in range(start_row, max(ws.max_row, start_row) + 1):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None


def _all_quality_rows(contract: dict[str, Any]) -> list[tuple[str, str | None, dict[str, Any]]]:
    quality_rows: list[tuple[str, str | None, dict[str, Any]]] = []
    for schema in contract.get("schema") or []:
        schema_name = schema.get("name") or schema.get("id")
        for quality in schema.get("quality") or []:
            quality_rows.append((schema_name, _schema_quality_property(schema, quality), quality))
        for prop in schema.get("properties") or []:
            for quality in prop.get("quality") or []:
                quality_rows.append((schema_name, prop.get("name") or prop.get("id"), quality))
    return quality_rows


def _schema_quality_property(schema: dict[str, Any], quality: dict[str, Any]) -> str | None:
    if quality.get("metric") == "rowCount":
        return None

    for item in quality.get("customProperties") or []:
        if item.get("property") == "referencedColumns":
            return _join_values(item.get("value"))

    description = str(quality.get("description") or "")
    match = re.search(r"-\s+(.+?)\s+(?:primary key|business key)\b", description, flags=re.IGNORECASE)
    if match:
        return match.group(1).strip()

    query = str(quality.get("query") or "")
    match = re.search(r"concat\(([^)]+)\)", query, flags=re.IGNORECASE)
    if match:
        return ", ".join(part.strip() for part in match.group(1).split(",") if part.strip())

    return None


def _write_fundamentals(wb: Any, contract: dict[str, Any]) -> None:
    ws = wb["Fundamentals"]
    ws["C4"] = contract.get("kind")
    ws["C5"] = contract.get("apiVersion")
    ws["C7"] = contract.get("id")
    ws["C8"] = contract.get("name")
    ws["C9"] = contract.get("version")
    ws["C10"] = contract.get("status")

    team = contract.get("team") or {}
    members = team.get("members") or []
    owner = members[0].get("username") if members else None
    ws["C12"] = owner or team.get("name")
    ws["C14"] = contract.get("domain")
    ws["C15"] = contract.get("dataProduct")
    ws["C16"] = contract.get("tenant")

    description = contract.get("description") or {}
    if isinstance(description, dict):
        ws["C19"] = description.get("purpose")
        ws["C20"] = description.get("limitations")
        ws["C21"] = description.get("usage")
    else:
        ws["C19"] = description
    ws["C23"] = _join_values(contract.get("tags"))


def _write_schema_sheets(wb: Any, contract: dict[str, Any]) -> list[str]:
    schema_template = wb["Schema <table_name>"]
    used_titles = set(wb.sheetnames)
    created_sheets: list[str] = []
    property_cols = {
        "name": 1,
        "businessName": 2,
        "logicalType": 3,
        "physicalType": 4,
        "examples": 5,
        "description": 6,
        "required": 7,
        "unique": 8,
        "classification": 9,
        "tags": 10,
        "authoritativeDefinitionUrl": 11,
        "authoritativeDefinitionType": 12,
        "physicalName": 13,
        "primaryKey": 14,
        "primaryKeyPosition": 15,
        "partitioned": 16,
        "partitionKeyPosition": 17,
        "encryptedName": 18,
        "transformSources": 19,
        "transformLogic": 20,
        "transformDescription": 21,
        "criticalDataElement": 22,
        "maxItems": 23,
        "minItems": 24,
        "uniqueItems": 25,
        "format": 26,
        "minLength": 27,
        "maxLength": 28,
        "exclusiveMinimum": 29,
        "minimum": 30,
        "exclusiveMaximum": 31,
        "maximum": 32,
        "multipleOf": 33,
        "minProperties": 34,
        "maxProperties": 35,
        "requiredProperties": 36,
        "pattern": 37,
    }

    for schema in contract.get("schema") or []:
        table_name = schema.get("name") or schema.get("id")
        sheet = wb.copy_worksheet(schema_template)
        sheet.title = _sheet_title(str(table_name), used_titles)
        _add_schema_defined_names(sheet)
        created_sheets.append(sheet.title)

        sheet["A1"] = f"Schema {table_name}"
        sheet["B5"] = table_name
        sheet["B6"] = schema.get("physicalType")
        sheet["B7"] = schema.get("description")
        sheet["B8"] = schema.get("businessName")
        sheet["B9"] = schema.get("physicalName")
        sheet["B10"] = schema.get("dataGranularityDescription")
        sheet["B11"] = _join_values(schema.get("tags"))

        for row_index, prop in enumerate(schema.get("properties") or [], start=14):
            if row_index > sheet.max_row:
                _copy_row_style(sheet, 14, row_index, 38)
            for key, col in property_cols.items():
                value = prop.get(key)
                if key in {"examples", "tags", "transformSources", "requiredProperties"}:
                    value = _join_values(value)
                sheet.cell(row_index, col).value = _scalar(value)

            comments: list[dict[str, Any]] = []
            if prop.get("id") and prop.get("id") != prop.get("name"):
                comments.append({"id": prop.get("id")})
            if prop.get("customProperties"):
                comments.append({"customProperties": prop.get("customProperties")})
            if comments:
                sheet.cell(row_index, 38).value = json.dumps(comments, ensure_ascii=False)

    wb.remove(schema_template)
    return created_sheets


def _write_quality(wb: Any, contract: dict[str, Any]) -> int:
    ws = wb["Quality"]
    _clear_rows(ws, 5, 13)
    quality_rows = _all_quality_rows(contract)

    for offset, (schema_name, prop_name, quality) in enumerate(quality_rows):
        row = 5 + offset
        if row > ws.max_row:
            _copy_row_style(ws, 5, row, 13)
        operator, operator_value = _quality_operator(quality)
        excel_threshold_value = operator_value
        if isinstance(operator_value, (dict, list)) and operator != "mustBeBetween":
            excel_threshold_value = _first_numeric_value(operator_value)
        rule_label = _rule_label(quality)
        ws.cell(row, 1).value = schema_name
        ws.cell(row, 2).value = prop_name
        ws.cell(row, 3).value = "custom"
        ws.cell(row, 4).value = quality.get("description")
        ws.cell(row, 5).value = None
        ws.cell(row, 6).value = None
        ws.cell(row, 7).value = operator
        ws.cell(row, 8).value = _scalar(excel_threshold_value)
        ws.cell(row, 9).value = "business_rule_catalog"
        ws.cell(row, 10).value = rule_label
        ws.cell(row, 11).value = quality.get("severity")
        ws.cell(row, 12).value = quality.get("scheduler")
        ws.cell(row, 13).value = quality.get("schedule")
    return len(quality_rows)


def _write_list_sheets(wb: Any, contract: dict[str, Any]) -> None:
    support_ws = wb["Support"]
    _clear_rows(support_ws, 5, 6)
    for offset, item in enumerate(contract.get("support") or []):
        row = 5 + offset
        if row > support_ws.max_row:
            _copy_row_style(support_ws, 5, row, 6)
        support_ws.cell(row, 1).value = item.get("channel")
        support_ws.cell(row, 2).value = item.get("url")
        support_ws.cell(row, 3).value = item.get("description")
        support_ws.cell(row, 4).value = item.get("tool")
        support_ws.cell(row, 5).value = item.get("scope")
        support_ws.cell(row, 6).value = item.get("invitationUrl")

    team = contract.get("team") or {}
    members = team.get("members") or []
    team_ws = wb["Team"]
    _clear_rows(team_ws, 5, 7)
    for offset, member in enumerate(members):
        row = 5 + offset
        if row > team_ws.max_row:
            _copy_row_style(team_ws, 5, row, 7)
        team_ws.cell(row, 1).value = member.get("username")
        team_ws.cell(row, 2).value = member.get("name") or member.get("username")
        team_ws.cell(row, 3).value = member.get("description") or team.get("description")
        team_ws.cell(row, 4).value = member.get("role")
        team_ws.cell(row, 5).value = member.get("dateIn")
        team_ws.cell(row, 6).value = member.get("dateOut")
        team_ws.cell(row, 7).value = member.get("replacedByUsername")

    roles_ws = wb["Roles"]
    _clear_rows(roles_ws, 5, 5)
    for offset, role in enumerate(contract.get("roles") or []):
        row = 5 + offset
        if row > roles_ws.max_row:
            _copy_row_style(roles_ws, 5, row, 5)
        roles_ws.cell(row, 1).value = role.get("role")
        roles_ws.cell(row, 2).value = role.get("description")
        roles_ws.cell(row, 3).value = role.get("access")
        roles_ws.cell(row, 4).value = role.get("firstLevelApprovers")
        roles_ws.cell(row, 5).value = role.get("secondLevelApprovers")

    sla_ws = wb["SLA"]
    _clear_rows(sla_ws, 7, 6)
    for offset, sla in enumerate(contract.get("slaProperties") or []):
        row = 7 + offset
        if row > sla_ws.max_row:
            _copy_row_style(sla_ws, 7, row, 6)
        sla_ws.cell(row, 1).value = sla.get("property")
        sla_ws.cell(row, 2).value = _scalar(sla.get("value"))
        sla_ws.cell(row, 3).value = _scalar(sla.get("extendedValue"))
        sla_ws.cell(row, 4).value = sla.get("unit")
        sla_ws.cell(row, 5).value = sla.get("element")
        sla_ws.cell(row, 6).value = sla.get("driver")


def _write_servers_and_custom_properties(wb: Any, contract: dict[str, Any]) -> None:
    servers = contract.get("servers") or []
    if servers:
        server = servers[0]
        ws = wb["Servers"]
        ws["B4"] = server.get("server")
        ws["B5"] = server.get("environment")
        ws["B6"] = server.get("description")
        ws["B8"] = server.get("type")
        ws["B12"] = server.get("path") or server.get("location")
        ws["B13"] = server.get("format")
        ws["B14"] = server.get("delimiter")

    pricing = contract.get("pricing") or {}
    if pricing:
        ws = wb["Pricing"]
        ws["B4"] = pricing.get("priceAmount")
        ws["B5"] = pricing.get("priceCurrency")
        ws["B6"] = pricing.get("priceUnit")

    ws = wb["Custom Properties"]
    _clear_rows(ws, 5, 2)
    custom_rows: list[tuple[str | None, Any]] = [
        (item.get("property"), item.get("value")) for item in contract.get("customProperties") or []
    ]
    for item in contract.get("authoritativeDefinitions") or []:
        custom_rows.append(("authoritativeDefinition", item))
    if servers:
        custom_rows.append(("servers", servers))
    custom_rows.append(("contractCreatedTs", contract.get("contractCreatedTs")))

    for offset, (prop, value) in enumerate(custom_rows):
        row = 5 + offset
        if row > ws.max_row:
            _copy_row_style(ws, 5, row, 2)
        ws.cell(row, 1).value = prop
        ws.cell(row, 2).value = _scalar(value)


def export_contract_to_excel(
    contract_path: str | Path,
    output_path: str | Path,
    template_path: str | Path = DEFAULT_EXCEL_TEMPLATE,
) -> dict[str, Any]:
    contract_file = Path(contract_path)
    template_file = Path(template_path)
    output_file = Path(output_path)

    if not template_file.exists():
        raise FileNotFoundError(f"ODCS Excel template not found: {template_file}")

    with contract_file.open(encoding="utf-8") as handle:
        contract = yaml.safe_load(handle) or {}

    wb = load_workbook(template_file)
    _write_fundamentals(wb, contract)
    created_schema_sheets = _write_schema_sheets(wb, contract)
    quality_count = _write_quality(wb, contract)
    _write_list_sheets(wb, contract)
    _write_servers_and_custom_properties(wb, contract)

    preferred_order = ["Instructions", "Fundamentals"] + created_schema_sheets + [
        "Relationships",
        "Quality",
        "Support",
        "Team",
        "Roles",
        "SLA",
        "Servers",
        "Pricing",
        "Custom Properties",
    ]
    wb._sheets = [wb[name] for name in preferred_order if name in wb.sheetnames]

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    return {
        "excel": str(output_file),
        "template": str(template_file),
        "schemaSheetCount": len(created_schema_sheets),
        "propertyCount": sum(len(schema.get("properties") or []) for schema in contract.get("schema") or []),
        "qualityCount": quality_count,
    }
