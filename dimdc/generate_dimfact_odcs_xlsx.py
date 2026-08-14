from __future__ import annotations

import argparse
import re
from copy import copy
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName


DEFAULT_TEMPLATE = Path("dimdc/odcs-template-updated.xlsx")
DEFAULT_TABLE_METADATA = Path("dimdc/dimfacttablemetadat.csv")
DEFAULT_RULES = Path("dimdc/rules.csv")
DEFAULT_RULE_MAPPING = Path("dimdc/dimfact_rule_mapping.csv")
DEFAULT_OUTPUT = Path("dimdc/outputs/dimfact_odcs_contract.xlsx")


PROPERTY_COLUMNS = [
    "Property",
    "Business Name",
    "Logical Type",
    "Physical Type",
    "Example(s)",
    "Description",
    "Required",
    "Unique",
    "Classification",
    "Tags",
    "Authoritative Definition URL",
    "Authoritative Definition Type",
]


QUALITY_COLUMNS = [
    "Schema",
    "Property",
    "Quality Type",
    "Description",
    "Rule (Library)",
    "Query (SQL)",
    "Threshold Operator",
    "Threshold Value",
    "Quality Engine (Custom)",
    "Implementation (Custom)",
    "Severity",
    "Scheduler",
    "Schedule",
]


RELATIONSHIP_COLUMNS = [
    "Level",
    "Type",
    "From",
    "To",
    "Description",
]


def _truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"true", "t", "1", "yes", "y"}


def _clean(value: Any) -> str:
    text = "" if pd.isna(value) else str(value).strip()
    return "" if text.lower() in {"nan", "none", "null"} else text


def _split_columns(value: Any) -> list[str]:
    return [part.strip() for part in _clean(value).split(",") if part.strip()]


def _business_name(value: str) -> str:
    return value.replace("_", " ").title()


def _sheet_title(table_name: str, used: set[str]) -> str:
    base = "Schema " + re.sub(r"[\\/*?:\[\]]", "_", table_name)
    if len(base) <= 31 and base not in used:
        used.add(base)
        return base
    for index in range(1, 1000):
        suffix = f"_{index}"
        title = base[: 31 - len(suffix)] + suffix
        if title not in used:
            used.add(title)
            return title
    raise RuntimeError(f"Unable to create unique sheet title for {table_name}")


def _copy_row_style(ws: Any, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def _quoted_sheet_title(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _add_local_defined_name(ws: Any, name: str, cell_range: str) -> None:
    ws.defined_names.add(DefinedName(name=name, attr_text=f"{_quoted_sheet_title(ws.title)}!{cell_range}"))


def _add_schema_defined_names(ws: Any) -> None:
    ws.defined_names.clear()
    _add_local_defined_name(ws, "schema.name", "$B$5")
    _add_local_defined_name(ws, "schema.physicalType", "$B$6")
    _add_local_defined_name(ws, "schema.description", "$B$7")
    _add_local_defined_name(ws, "schema.businessName", "$B$8")
    _add_local_defined_name(ws, "schema.physicalName", "$B$9")
    _add_local_defined_name(ws, "schema.dataGranularityDescription", "$B$10")
    _add_local_defined_name(ws, "schema.tags", "$B$11")
    _add_local_defined_name(ws, "schema.properties", "$A$13:$AZ$981")


def _find_row(ws: Any, label: str) -> int:
    normalized = label.strip().lower()
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, str) and value.strip().lower() == normalized:
                return row
    raise RuntimeError(f"Could not find row with label {label!r} in sheet {ws.title!r}")


def _set_label_value(ws: Any, label: str, value: Any) -> None:
    normalized = label.strip().lower()
    for row in range(1, ws.max_row + 1):
        for col in range(1, min(ws.max_column, 6) + 1):
            cell_value = ws.cell(row, col).value
            if isinstance(cell_value, str) and cell_value.strip().lower() == normalized:
                ws.cell(row, col + 1).value = value
                return


def _clear_rows(ws: Any, start_row: int, max_col: int) -> None:
    for row in range(start_row, ws.max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None


def _property_rows(table: dict[str, Any]) -> list[dict[str, Any]]:
    pk = _clean(table["table_pk"])
    bk_columns = _split_columns(table["table_bk"])
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()

    def add_property(
        name: str,
        logical_type: str,
        physical_type: str,
        description: str,
        required: bool = True,
        unique: bool = False,
        classification: str = "internal",
        tags: str = "",
    ) -> None:
        if not name or name in seen:
            return
        seen.add(name)
        rows.append({
            "Property": name,
            "Business Name": _business_name(name),
            "Logical Type": logical_type,
            "Physical Type": physical_type,
            "Example(s)": "",
            "Description": description,
            "Required": "TRUE" if required else "FALSE",
            "Unique": "TRUE" if unique else "FALSE",
            "Classification": classification,
            "Tags": tags,
            "Authoritative Definition URL": "",
            "Authoritative Definition Type": "",
        })

    add_property(
        pk,
        "string",
        "varchar(64)",
        "Primary key for this gold table.",
        required=True,
        unique=True,
        classification="restricted" if pk.endswith("_sk") else "internal",
        tags="primary_key",
    )

    composite_bk = len(bk_columns) > 1
    for column in bk_columns:
        add_property(
            column,
            "string",
            "varchar(128)" if composite_bk else "varchar(64)",
            "Business key for this gold table.",
            required=True,
            unique=not composite_bk and column != pk,
            classification="confidential",
            tags="business_key",
        )

    if _truthy(table["scd2"]):
        add_property("effective_from_ts", "timestamp", "timestamp", "SCD2 effective start timestamp.", tags="scd2")
        add_property("effective_to_ts", "timestamp", "timestamp", "SCD2 effective end timestamp.", tags="scd2")
        add_property("is_current", "boolean", "boolean", "Current active record indicator.", tags="scd2")

    add_property("load_ts", "timestamp", "timestamp", "Gold table load timestamp.", required=False, tags="audit")
    add_property("record_source", "string", "varchar(128)", "Source system or source table for the row.", required=False, tags="audit")
    return rows


def _quality_property(rule_id: str, table: dict[str, Any]) -> str:
    if rule_id == "rule_0003":
        return _clean(table["table_pk"])
    if rule_id in {"rule_0002", "rule_0004", "rule_0005", "rule_0006"}:
        return ""
    if rule_id in {"rule_0001", "rule_0002", "rule_0005", "rule_0006", "rule_0007"}:
        table_bk = _clean(table["table_bk"])
        return "" if "," in table_bk else table_bk
    if rule_id == "rule_0004":
        return "effective_from_ts, effective_to_ts"
    return _clean(table["table_bk"] or table["table_pk"])


def _write_fundamentals(wb: Any) -> None:
    if "Fundamentals" not in wb.sheetnames:
        return
    ws = wb["Fundamentals"]
    values = {
        "Kind": "DataContract",
        "API Version": "v3.1.0",
        "ID": "dimfact-gold-odcs-contract",
        "Name": "Dim Fact Gold ODCS Data Contract",
        "Version": "1.0.0",
        "Status": "active",
        "Domain": "insurance",
        "Data Product": "business_vault_gold",
        "Tenant": "allianz",
        "Purpose": "Data contract for gold dimension and fact tables generated from approved dimdc metadata and rule mappings.",
        "Limitations": "Column coverage is limited to table metadata fields supplied in dimfacttablemetadat.csv.",
        "Usage": "Used to document table keys, source lineage, and active data quality rule assignments.",
    }
    for label, value in values.items():
        _set_label_value(ws, label, value)


def _write_schema_sheet(ws: Any, table: dict[str, Any]) -> None:
    _add_schema_defined_names(ws)
    table_name = _clean(table["table_name"])
    table_bk = _clean(table["table_bk"])
    source_table = _clean(table["source_table"])

    _set_label_value(ws, "Name", table_name)
    _set_label_value(ws, "Type", "table")
    _set_label_value(
        ws,
        "Description",
        f"Gold {table_name} table generated from dim/fact metadata.",
    )
    _set_label_value(ws, "Business Name", _business_name(table_name))
    _set_label_value(ws, "Physical Name", table_name)
    _set_label_value(ws, "Data Granularity", f"One row per {table_bk or table['table_pk']}.")
    _set_label_value(ws, "Tags", f"{table['stage']}, {table['schema_name']}, {'scd2' if _truthy(table['scd2']) else 'snapshot'}")

    header_row = _find_row(ws, "Property")
    data_start = header_row + 1
    _clear_rows(ws, data_start, len(PROPERTY_COLUMNS))
    for index, row in enumerate(_property_rows(table), start=data_start):
        _copy_row_style(ws, data_start, index, len(PROPERTY_COLUMNS))
        for col, key in enumerate(PROPERTY_COLUMNS, start=1):
            ws.cell(index, col).value = row.get(key)

    if source_table:
        ws.cell(data_start + len(_property_rows(table)) + 1, 1).value = "Source Table"
        ws.cell(data_start + len(_property_rows(table)) + 1, 2).value = source_table


def _write_relationships(wb: Any, tables: list[dict[str, Any]]) -> None:
    if "Relationships" not in wb.sheetnames:
        return
    ws = wb["Relationships"]
    header_row = _find_row(ws, "Level")
    data_start = header_row + 1
    _clear_rows(ws, data_start, len(RELATIONSHIP_COLUMNS))

    row_number = data_start
    for table in tables:
        source_table = _clean(table["source_table"])
        source_pk = _clean(table["source_pk"])
        target_bk = _clean(table["table_bk"])
        table_name = _clean(table["table_name"])
        if not source_table or not source_pk or not target_bk:
            continue
        _copy_row_style(ws, data_start, row_number, len(RELATIONSHIP_COLUMNS))
        values = {
            "Level": "property",
            "Type": "lineage",
            "From": f"{source_table}.{source_pk}",
            "To": f"{table_name}.{target_bk}",
            "Description": "Source-to-gold key lineage from dim/fact metadata.",
        }
        for col, key in enumerate(RELATIONSHIP_COLUMNS, start=1):
            ws.cell(row_number, col).value = values[key]
        row_number += 1


def _write_quality(wb: Any, tables: list[dict[str, Any]], rules: pd.DataFrame, mappings: pd.DataFrame) -> None:
    if "Quality" not in wb.sheetnames:
        return
    ws = wb["Quality"]
    header_row = _find_row(ws, "Schema")
    data_start = header_row + 1
    _clear_rows(ws, data_start, len(QUALITY_COLUMNS))

    active_tables = {table["table_id"]: table for table in tables if _truthy(table["is_active"])}
    active_rules = {
        row["rule_id"]: row
        for row in rules.to_dict("records")
    }

    row_number = data_start
    for mapping in mappings.to_dict("records"):
        if not _truthy(mapping["is_active"]):
            continue
        table = active_tables.get(mapping["table_id"])
        rule = active_rules.get(mapping["rule_id"])
        if not table or not rule:
            continue

        rule_id = _clean(rule["rule_id"])
        rule_name = _clean(rule["rule_name"])
        severity = _clean(rule["rule_type"]) or "Warning"
        description = _clean(rule["query_desc"]) or rule_name
        _copy_row_style(ws, data_start, row_number, len(QUALITY_COLUMNS))
        values = {
            "Schema": _clean(table["table_name"]),
            "Property": _quality_property(rule_id, table),
            "Quality Type": "custom",
            "Description": f"{rule_id}: {rule_name} - {description}",
            "Rule (Library)": rule_id,
            "Query (SQL)": "",
            "Threshold Operator": "",
            "Threshold Value": "",
            "Quality Engine (Custom)": "business_rule_catalog",
            "Implementation (Custom)": f"{rule_id}: {rule_name}",
            "Severity": severity,
            "Scheduler": "",
            "Schedule": "",
        }
        for col, key in enumerate(QUALITY_COLUMNS, start=1):
            ws.cell(row_number, col).value = values[key]
        row_number += 1


def generate_dimfact_odcs_xlsx(
    template: Path,
    table_metadata_path: Path,
    rules_path: Path,
    rule_mapping_path: Path,
    output_path: Path,
) -> Path:
    table_metadata = pd.read_csv(table_metadata_path).fillna("")
    rules = pd.read_csv(rules_path).fillna("")
    mappings = pd.read_csv(rule_mapping_path).fillna("")

    active_tables = [
        row for row in table_metadata.to_dict("records")
        if _truthy(row["is_active"])
    ]

    wb = load_workbook(template)
    _write_fundamentals(wb)

    schema_templates = [name for name in wb.sheetnames if name.lower().startswith("schema")]
    if not schema_templates:
        raise RuntimeError("Template does not contain a Schema sheet to copy.")
    template_ws = wb[schema_templates[0]]
    used_titles = {name for name in wb.sheetnames if not name.lower().startswith("schema")}

    created_sheets = []
    for table in active_tables:
        ws = wb.copy_worksheet(template_ws)
        ws.title = _sheet_title(_clean(table["table_name"]), used_titles)
        _write_schema_sheet(ws, table)
        created_sheets.append(ws.title)

    for name in schema_templates:
        if name in wb.sheetnames:
            del wb[name]

    fundamentals_index = wb.sheetnames.index("Fundamentals") + 1 if "Fundamentals" in wb.sheetnames else 0
    for offset, name in enumerate(created_sheets):
        wb.move_sheet(wb[name], offset=-(wb.sheetnames.index(name) - fundamentals_index - offset))

    _write_relationships(wb, active_tables)
    _write_quality(wb, active_tables, rules, mappings)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate an ODCS Excel workbook for dim/fact tables.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--table-metadata", default=str(DEFAULT_TABLE_METADATA))
    parser.add_argument("--rules", default=str(DEFAULT_RULES))
    parser.add_argument("--rule-mapping", default=str(DEFAULT_RULE_MAPPING))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()

    output = generate_dimfact_odcs_xlsx(
        template=Path(args.template),
        table_metadata_path=Path(args.table_metadata),
        rules_path=Path(args.rules),
        rule_mapping_path=Path(args.rule_mapping),
        output_path=Path(args.output),
    )
    print(f"Generated: {output}")


if __name__ == "__main__":
    main()
